# -*- coding: utf-8 -*-
"""Framework v2.14 — the workbook edit for the V6.0 production round, second part (owner decisions, 22 Sep 2026).

    python -m pipeline.make_framework_v214 config/01_Framework_v2.13.xlsx config/01_Framework_v2.14.xlsx

  sheet 49  EN-11: for a school with a year group inside its range that has no
            accepted response (58 schools in the register), the overview sentence
            keeps the range and drops the word "All" — owner instruction of
            22 Sep 2026 ("keep the range and drop the word, change nothing else").
  sheet 43  ONE new frame, ``ui.overview_note_gap``: the same sentence as
            ui.overview_note without "All"; the client selects it when the
            payload's school block carries yearsGap. Welsh EMPTY -> TRANSLATOR.
  sheet 74  change log; sheet 75 flags (the four sheet-53 local-authority rows the
            translator is asked for under option A). No Welsh composed.
"""
from __future__ import annotations

import datetime
import sys

import openpyxl

KEY = "ui.overview_note_gap"
EN = ("Year groups from Year {first} to Year {last} are represented. {hi} contributed the largest number of "
      "responses ({hin}) and {lo} the smallest ({lon}). Every response option is shown, including those selected "
      "by no pupils. Unequal group sizes should be kept in mind when comparing raw counts.")


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    assert "74 v2.14 change log" not in wb.sheetnames
    ws43 = wb["43 Interface frames"]
    keys = {str(r[0].value) for r in ws43.iter_rows(min_row=4) if r[0].value}
    assert KEY not in keys
    ws43.append([KEY, "V6.0 EN-11", "#overview-note — used instead of ui.overview_note when a year group inside the school's range has no accepted response (payload school.yearsGap)",
                 EN, None, "hi, hin, lo, lon, first, last",
                 "V6.0 (owner decision, 22 Sep 2026): ui.overview_note without the word \"All\" — the range is kept, nothing else changes. "
                 "Welsh pending the translator (the ui.overview_note Welsh minus \"holl\" is for the translator to confirm, not composed here).",
                 "TRANSLATOR"])
    h = ws43.cell(row=1, column=1)
    h.value = str(h.value) + " · v2.14: +ui.overview_note_gap (EN-11; Welsh pending) — see sheet 74."
    ws49 = wb["49 English client edits"]
    ws49.append(["EN-11", "ui.overview_note for a school with a year group inside its range that has no accepted response",
                 "All year groups from Year {first} to Year {last} are represented.",
                 "Year groups from Year {first} to Year {last} are represented. (frame ui.overview_note_gap; every other word unchanged)",
                 "Owner instruction of 22 Sep 2026: keep the range and drop the word, change nothing else. 58 schools in the 22 Sep register. "
                 "No narrative state changes; the frame is client-realised.",
                 "SANCTIONED v2.14 (owner, 22 Sep 2026)"])
    ws74 = wb.create_sheet("74 v2.14 change log")
    ws74.append([f"Framework v2.14 change log — generated {datetime.date.today().isoformat()} by pipeline/make_framework_v214.py from v2.13. "
                 "Sheet 49 EN-11 (owner-sanctioned); sheet 43 +ui.overview_note_gap. No Welsh composed; no translator wording changed; no narrative surface touched."])
    ws74.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    ws74.append(["43 Interface frames", KEY, "NEW ROW", "", EN, "EN-11 — owner instruction, 22 Sep 2026"])
    ws74.append(["49 English client edits", "EN-11", "NEW ROW", "", "range kept, 'All' dropped for gap-year schools", "owner instruction, 22 Sep 2026"])
    ws75 = wb.create_sheet("75 V6.0 flags (2)")
    ws75.append(["Flags raised while applying the owner's decisions of 22 Sep 2026 — questions for the translator; nothing here is decided by Industryline."])
    ws75.append(["Where", "What", "Detail", "Action needed"])
    ws75.append(["43 Interface frames — ui.overview_note_gap", "Welsh for the gap-year variant",
                 "English is ui.overview_note without 'All'. The corresponding Welsh is for the translator (the existing Welsh minus 'holl' is the obvious reading but is not composed here).",
                 "Translator: one row in the next handoff."])
    ws75.append(["53 Proper names — four local authorities (option A)", "Rows in the dataset's spelling",
                 "The dataset names the authorities 'Carmarthenshire' (no row), 'Conwy' (sheet: Conwy County), 'Rhondda Cynon Taf' (sheet: Rhondda Borough — check), 'The Vale of Glamorgan' (sheet: Vale of Glamorgan). "
                 "211 schools are HELD by the register until the rows exist.",
                 "Translator: four sheet-53 rows, English exactly as the dataset spells it, Welsh as ruled."])
    ws0 = wb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = ("Version 2.14 (V6.0, 22 Sep 2026): sheet 49 EN-11 (gap-year overview sentence: range kept, 'All' dropped — owner); "
                                       "sheet 43 +ui.overview_note_gap (Welsh pending); sheet 74 change log; sheet 75 flags (four sheet-53 rows requested). No translator wording changed.")
    wb.save(dst)
    print(f"written {dst}")


if __name__ == "__main__":
    main()

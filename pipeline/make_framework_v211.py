# -*- coding: utf-8 -*-
"""Framework v2.11 — the workbook edit for V4.18 (owner instruction, 22 Sep 2026).

    python -m pipeline.make_framework_v211 config/01_Framework_v2.10.xlsx config/01_Framework_v2.11.xlsx

  sheet 49  EN-09: two sanctioned changes to the locked narrative/report —
            (a) the Club Sports section (module d2, the club-sport square-root
            estimate chart, its narrative and its appendix table) is removed
            from the report; (b) under a SETTING selection ("Where are our
            pupils taking part in Sport?") the weekly-frequency chart keeps the
            wider picture and its narrative leads with the existing
            selected-group definition sentence (template group_defined_v2,
            unchanged wording).
  sheet 43  ui.appx_club → DEPRECATED (its consumer, the appendix table, is
            gone; FRAME-dormant reads the status).
  sheet 68  change log; sheet 69 V4.18 flags. No Welsh composed; no
            translator wording changed.
"""
from __future__ import annotations

import datetime
import sys

import openpyxl


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    assert "68 v2.11 change log" not in wb.sheetnames
    log = []
    ws43 = wb["43 Interface frames"]
    rows43 = {str(r[0].value): r for r in ws43.iter_rows(min_row=4) if r[0].value}
    r = rows43["ui.appx_club"]
    old_status = str(r[7].value or "")
    r[7].value = "DEPRECATED"
    r[6].value = ((r[6].value or "") + " DEPRECATED v2.11 (V4.18): the Club Sports section and its appendix table "
                  "were removed from the report (owner instruction, 22 Sep 2026); no consumer.").strip()
    log.append(("43 Interface frames", "ui.appx_club", "DEPRECATED", old_status, "DEPRECATED",
                "V4.18 — Club Sports section removed (owner instruction, 22 Sep 2026)"))
    h = ws43.cell(row=1, column=1)
    h.value = str(h.value) + " · v2.11: ui.appx_club DEPRECATED — see sheet 68."
    ws49 = wb["49 English client edits"]
    ws49.append(["EN-09", "report structure + narrative module d0 (freq_estimate) under a participation_settings selection",
                 "Club Sports section (h3 ‘Club Sports’, its context paragraph, the club-sport weekly-estimate chart, module d2, appendix table ‘Club sport (school or community)’) present; "
                 "under a setting selection the weekly-frequency chart showed the selected group's own distribution and its narrative had no definition sentence.",
                 "Club Sports section removed from the report (module d2 no longer generated; club_freq_estimate no longer a selectable source; the metric is still computed for the summary rows that cite it). "
                 "Under a setting selection the weekly-frequency chart keeps the wider picture (the demographic view's bars, still selectable) and its narrative leads with the existing definition sentence "
                 "‘This selected group contains {phrase}. The chart above keeps the wider picture for {demo} for context, with the selected answer highlighted.’ (template group_defined_v2, wording unchanged), "
                 "followed by the group's sentences as before.",
                 "Owner instruction of 22 Sep 2026. The English corpus moves by this ruling: every state loses its d2 paragraphs, the cb_* selected-group states (10 × 36) disappear, and each st_* state's d0 gains one definition paragraph; "
                 "every other paragraph is byte-identical to V4.17 (state-by-state proof in the V4.18 record). Lock re-emitted as 02c_lock_V418_v8.json.",
                 "SANCTIONED v2.11 (owner, 22 Sep 2026)"])
    log.append(("49 English client edits", "EN-09", "NEW ROW", "", "Club Sports removed; d0 definition sentence under a setting selection", "owner instruction, 22 Sep 2026"))
    ws68 = wb.create_sheet("68 v2.11 change log")
    ws68.append([f"Framework v2.11 change log — generated {datetime.date.today().isoformat()} by pipeline/make_framework_v211.py from v2.10. "
                 "Sheet 49 EN-09 (owner-sanctioned report change); sheet 43 ui.appx_club deprecated. No Welsh composed; no translator wording changed."])
    ws68.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    for row in log:
        ws68.append(list(row))
    ws69 = wb.create_sheet("69 V4.18 flags")
    ws69.append(["Flags raised while preparing V4.18 — each is a question for the translator, Sport Wales or the owner; nothing here is decided by Industryline."])
    ws69.append(["Where", "What", "Detail", "Action needed"])
    ws69.append(["Report — summary page and chapter summary", "The club-sport estimate still appears in two places",
                 "The Club Sports section is gone, but the summary table row ‘Club sport 3+ times a week (estimated)’ / ‘Less than weekly club sport’ and the chapter-summary sentence ‘… took part in club sport at least once a week’ are computed from the same metric and were left as they were (locked narrative).",
                 "Owner: confirm they stay, or sanction their removal (EN-10)."])
    ws69.append(["Translator handoff", "Two static rows retired",
                 "The heading ‘Club Sports’ and its context paragraph are no longer on the page; the translator's Welsh for them is kept on the handoff's Retired sheet (V4.18 handoff).",
                 "None — recorded."])
    ws69.append(["Corpus lock", "Re-emitted under EN-09",
                 "02c_lock_V418_v8.json supersedes the V4.15 lock: d2 paragraphs removed everywhere, cb_* states removed, one definition paragraph added to d0 in every st_* state; all other paragraphs byte-identical to V4.17.",
                 "None — the V4.18 record carries the proof."])
    ws0 = wb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = ("Version 2.11 (V4.18, 22 Sep 2026): sheet 49 EN-09 — the Club Sports section removed and the weekly-frequency chart's "
                                       "definition sentence under a setting selection (owner-sanctioned); sheet 43 ui.appx_club DEPRECATED; sheet 68 change log; "
                                       "sheet 69 V4.18 flags. No translator wording changed.")
    wb.save(dst)
    print(f"written {dst}")


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""Apply the translator's completed handover (21 Sep 2026) to the workbooks.

    python -m pipeline.apply_translator_handover <completed_handover.xlsx> \\
        <framework_in.xlsx> <framework_out.xlsx> \\
        <handoff_in.xlsx> <handoff_out.xlsx> <register_out.xlsx> <translator.docx>

The translator's document (15 Sep 2026) is read for the one value the
handover quoted truncated (the ethnicity profile's White label, sheet 2b
row 49): the workbook cell is checked to be a prefix of the document's line
and the document's full line is used.

Reads the six sheets of SSS2026_Welsh_Translator_Handover_V4.14 as the
translator returned it and produces, deterministically:

  * the translation HANDOFF (Strings sheet) for V4.15 — sheet 1 answers
    (reviewer suggestions the translator confirmed, applied as the exact
    scripted substitutions of apply_review_corrections; the translator's own
    versions, whole cell), sheet 2 answers (the suggested wordings, as exact
    substitutions; the translator's version of ui009), sheet 3 (the page text
    that had no Welsh: every yellow cell, verbatim after the typographic
    normalisation of rule 41), and the AW-4 term decisions that touch page
    text (cyflwr hirdymor; "y wedd hon" → "y golwg hwn");
  * the FRAMEWORK v2.7 — sheet 43 frames from sheet 3, sheet 23 labels and
    the new profile rows from sheet 2b, sheet 58 chip prefixes from sheet 2b,
    sheets 22 / 31 / 44 / 17 for the long-term-condition term, sheet 11
    CONJ-06 (the figure rule, sheet 4 rule 15 / sheet 5 #12), sheet 09 (ail +
    soft mutation, sheet 4 rule 16 / sheet 5 #34), sheet 39 confirmations,
    sheet 02 decisions D84–D87, sheet 60 change log, sheet 61 flags;
  * a REGISTER workbook listing every applied change with its source row and
    every flag — a decision that contradicts another, a typo left verbatim, a
    decision applied by extension of its siblings, a decision not applied.

Rules honoured: the translator's text is applied word for word (typography
apart); no Welsh is composed here — every value is a cell of the handover,
a cell of the approved survey (framework sheet 30 / 23) or a cell that was
already in the workbook; a contradiction is decided by the more specific
answer (a per-row decision over a one-term-for-everything answer, an answer
with an explanatory comment over its dropdown) and ALWAYS flagged; a
substitution that does not match exactly once is an error, never skipped.
"""
from __future__ import annotations

import datetime
import json
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .apply_review_corrections import CORRECTIONS

TODAY = "2026-09-21"
SRC = "translator's completed handover, 21 Sep 2026"

# --------------------------------------------------------------- typography
def typo(s: str, log: list | None = None, where: str = "") -> str:
    """Rule 41 (sheet 4 rule 41, confirmed Correct): curly apostrophes and
    quotes throughout. Straight ASCII quotes typed in Excel are normalised;
    wording is never touched. Leading/trailing whitespace stripped."""
    if s is None:
        return s
    out = str(s).strip()
    before = out
    # non-breaking spaces pasted from Word become ordinary spaces (they
    # would forbid line-wrapping on the page); nothing else about spacing
    out = out.replace("\xa0", " ")
    # a straight apostrophe opening a quotation (after a space, a bracket or
    # at the start) is ‘; every other one is ’ (elision, possession, closing)
    out = re.sub(r"(^|[\s(\[])'", "\\1‘", out)
    out = out.replace("'", "’")
    # "…" pairs → “…”
    def _dq(m):
        return "“" + m.group(1) + "”"
    out = re.sub(r'"([^"]*)"', _dq, out)
    if log is not None and out != before:
        log.append({"where": where, "kind": "typography (rule 41)",
                    "before": before, "after": out})
    return out


# ---------------------------------------------------------------- readers
def read_sheet1(wb):
    ws = wb["1 Confirm corrections"]
    out = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        if not r[0] or not isinstance(r[0], str) or r[0].startswith(("Reference", "EXAMPLE")):
            continue
        key = r[0].split("\n")[0].strip()
        out[key] = {"answer": (r[6] or "").strip(), "version": r[7], "comment": r[8]}
    return out


def read_sheet2(wb):
    ws = wb["2 Decide wording"]
    rows, aw = {}, {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        if not r[0] or not isinstance(r[0], str):
            continue
        key = r[0].strip()
        if key.startswith("AW-"):
            aw[key] = {"answer": (r[6] or "").strip(), "version": r[7], "comment": r[8]}
        elif key not in ("Reference",) and not key.startswith("Answer options"):
            rows[key] = {"cy_now": r[3], "suggested": r[5],
                         "answer": (r[6] or "").strip(), "version": r[7], "comment": r[8]}
    return rows, aw


def read_sheet2b(wb):
    ws = wb["2b Answer labels"]
    rows = []
    for r in ws.iter_rows(min_row=7, values_only=True):
        if r[0] is None or not str(r[0]).strip().isdigit():
            continue
        rows.append(dict(n=int(r[0]), typ=r[1], pre_en=r[2], pre_today=r[3], pre_tr=r[5],
                         ans_en=r[6], ans_survey_en=r[7], ans_survey=r[8], ans_today=r[9],
                         ans_tr=r[11], dec_ans=(r[14] or "").strip(), dec_pre=(r[15] or "").strip(),
                         comment=r[16]))
    return rows


def read_sheet3(wb):
    ws = wb["3 Translate"]
    out = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        if not r[0] or not isinstance(r[0], str) or r[0] in ("Reference", "EXAMPLE"):
            continue
        if r[2] is None:            # section heading rows
            continue
        out[r[0].strip()] = {"en": r[2], "note": r[3], "cy": r[4], "comment": r[5]}
    return out


def read_sheet4(wb):
    ws = wb["4 Check the rules"]
    out = []
    for r in ws.iter_rows(min_row=7, values_only=True):
        if r[0] is None or not str(r[0]).strip().isdigit():
            continue
        out.append({"n": int(r[0]), "area": r[1], "rule": r[2], "ref": r[4],
                    "verdict": (r[5] or "").strip(), "comment": r[6]})
    return out


def read_sheet5(wb):
    ws = wb["5 Check sentences"]
    out = []
    for r in ws.iter_rows(min_row=7, values_only=True):
        if r[0] is None:
            continue
        k = str(r[0]).strip()
        if k.isdigit() or re.match(r"^T\d+$", k):
            out.append({"n": k, "verdict": (r[5] or "").strip(), "corrected": r[6], "why": r[7]})
    return out


# ------------------------------------------------- sheet 2 exact substitutions
# The suggestion column of sheet 2 is sometimes a fragment; the exact edit it
# stands for, against the V4.14 value, is fixed here (each must match once).
SHEET2_SUBSTITUTIONS = {
    "ui016": [("i ddeall beth sydd bwysicaf i blant a phobl ifanc yn well",
               "i ddeall yn well beth sydd bwysicaf i blant a phobl ifanc")],
    "ui084": [("Mae dau reswm pam nad yw canrannau efallai’n adio i 100%.",
               "Mae dau reswm pam efallai nad yw’r canrannau’n adio i 100%.")],
    "ui088": [("Yn ei thro, mae gwybodaeth yn cael ei chuddio",
               "O ganlyniad, mae gwybodaeth yn cael ei chuddio")],
    "ui099": [("pa mor aml mae disgyblion yn actif", "pa mor aml y mae disgyblion yn actif")],
    "ui109": [("yn yr Ysgol?", "yn yr ysgol?")],
    "ui110": [("ar draws 2 gyd-destun ysgol", "ar draws dau gyd-destun ysgol")],
    "ui120": [("y Tu allan i’r Ysgol?", "y tu allan i’r ysgol?")],
    "ui121": [("ar draws 2 gyd-destun y tu allan i’r ysgol", "ar draws dau gyd-destun y tu allan i’r ysgol")],
    "ui144": [("o gefndiroedd amrywiol yn ethnig", "o gefndiroedd ethnig amrywiol")],
    "ui163": [("ym mhob grŵp blwyddyn ddywedodd eu bod", "ym mhob grŵp blwyddyn a ddywedodd eu bod")],
    "ui166": [("Pa mor hyderus mae ein disgyblion", "Pa mor hyderus y mae ein disgyblion")],
    "ui202": [("Cliciwch i Neidio", "Cliciwch i neidio")],
    "ui217": [("deimlo: Yn iach?", "deimlo’n iach?")],
    "ui218": [("deimlo: Yn hyderus?", "deimlo’n hyderus?")],
    "ui219": [("deimlo: Yn barod i ddysgu?", "deimlo’n barod i ddysgu?")],
    "ui220": [("chwaraeon: mewn Gwersi AG?", "chwaraeon mewn gwersi AG?")],
    "ui221": [("chwaraeon: mewn Clybiau Chwaraeon Ysgol?", "chwaraeon mewn clybiau chwaraeon ysgol?")],
    "ui222": [("chwaraeon: mewn Clybiau’r tu allan i’r ysgol?", "chwaraeon mewn clybiau y tu allan i’r ysgol?")],
    "ui223": [("chwaraeon: Mewn Lleoliadau eraill, fel y parc neu’r ardd?",
               "chwaraeon mewn lleoliadau eraill, fel y parc neu’r ardd?")],
    "ui224": [("Beth mae’r Siart hon", "Beth mae’r siart hon")],
    "ui249": [("Mae disgyblion y dywedont eu bod", "Mae disgyblion a ddywedodd eu bod")],
    "ui256": [("yw bod y ddau wedi’u dewis", "yw bod y ddwy gamp wedi’u dewis")],
    "disability_condition": [("anabledd neu gyflwr tymor hir?", "anabledd neu gyflwr hirdymor?")],
    "settings_ed": [("o gefndiroedd amrywiol yn ethnig", "o gefndiroedd ethnig amrywiol")],
}
# sheet-43 frames decided on sheet 2
SHEET2_FRAME_SUBSTITUTIONS = {
    "ui.note_other_sports": [("unrhyw gampau a ddewisodd disgyblion", "unrhyw gamp a ddewisodd disgyblion")],
    "ui.stack_caption_sex": [],           # "rhywedd" confirmed — no edit
}

# --------------------------------------------------------------- sheet 3 map
# handover key -> where it lands
FRAME_KEYS_FROM_SHEET3 = [
    "ui.chart_sports_total_club", "ui.chart_sports_total_community", "ui.chart_sports_total_other",
    "ui.doc_title", "ui.fsm_context_title", "ui.fsm_context_body", "ui.fsm_context_value",
    "ui.review_note", "ui.meta_version_value", "ui.meta_survey_year", "ui.meta_report_version",
    "ui.meta_pipeline", "ui.meta_suppression", "ui.meta_generated", "ui.meta_checksum",
    "ui.meta_weighting", "ui.meta_weighting_value", "ui.gdpr_body_full", "ui.gdpr_body_module",
    "ui.a11y_switch", "ui.a11y_glossary_heading", "ui.a11y_glossary_intro",
    "ui.aria_cover", "ui.aria_toc", "ui.nav_aria",
]
NEW_HANDOFF_ROWS = {          # keys the payload declares but the Strings sheet never carried
    "h2_generic_reflect": ("discussion question", "Which findings match what you see day to day, and which are surprising?"),
    "h2_generic_discussion": ("discussion question", "Could further pupil discussion help explain these findings before provision changes are planned?"),
}

# --------------------------------------------------------- sheet 58 (2b → key)
PREFIX_EN_TO_FAM = {
    "Disability or long-term condition": "dy", "Ethnically Diverse Background": "ed",
    "Does sport": "st", "Estimated times active through sport each week": "ov",
    "Estimated times active through club sport each week": "cb", "Takes part": "tp",
    "Values": "im", "Would do more sport if": "mi", "Enjoys PE": "ep",
    "Enjoys School Sports Clubs": "eb", "Enjoys other settings": "eo", "Ideas listened to": "li",
    "Confidence trying a new sport": "ct", "Confidence learning a new skill": "cs",
    "Confidence trying again when sport is hard": "cg", "Ethnicity": "et",
}
LTC_OLD, LTC_NEW = "cyflwr tymor hir", "cyflwr hirdymor"
# the term appears mutated after neu / na / yng (gyflwr, chyflwr, nghyflwr)
LTC_RE = re.compile(r"\b(cyflwr|gyflwr|chyflwr|nghyflwr) tymor hir\b")


def has_ltc(text):
    return isinstance(text, str) and bool(LTC_RE.search(text))


def ltc_sub(text):
    return LTC_RE.sub(r"\1 hirdymor", text)


def sub_once(value, old, new, where):
    if value.count(old) != 1:
        raise SystemExit(f"{where}: expected exactly one occurrence of {old!r}; found {value.count(old)}")
    return value.replace(old, new)


class Register:
    def __init__(self):
        self.applied, self.flags, self.typo = [], [], []

    def add(self, target, key, change, before, after, source):
        self.applied.append({"target": target, "key": key, "change": change,
                             "before": before, "after": after, "source": source})

    def flag(self, key, what, detail, action):
        self.flags.append({"key": key, "what": what, "detail": detail, "action": action})


# ================================================================== HANDOFF
def apply_handoff(hwb, s1, s2, s3, aw, reg: Register):
    ws = hwb["Strings"]
    if (ws.cell(row=1, column=7).value or "") != "Note":
        ws.cell(row=1, column=7).value = "Note"
    rows = {str(r[3].value): r for r in ws.iter_rows(min_row=2) if r[3].value}

    def setv(key, new, note, change, source):
        r = rows[key]
        before = r[5].value
        r[5].value = new
        r[6].value = ((r[6].value or "") + " · " if r[6].value else "") + note
        reg.add("handoff", key, change, before, new, source)

    # ---- sheet 1 ---------------------------------------------------------
    for key, d in s1.items():
        if key not in rows:
            raise SystemExit(f"sheet 1: {key} not in the handoff")
        ans = d["answer"]
        cur = rows[key][5].value or ""
        if ans.startswith("Use the reviewer"):
            val = cur
            for old, new in CORRECTIONS[key]:
                val = sub_once(val, old, new, f"sheet 1 {key}")
            setv(key, val, f"translator confirmed the reviewer's suggestion ({SRC}, sheet 1)",
                 "reviewer suggestion confirmed → applied as the exact substitution", f"sheet 1 {key}: {ans}")
        elif ans.startswith("Use my version"):
            ver = typo(d["version"], reg.typo, f"sheet 1 {key}")
            if key == "ui250":
                m = re.search(r"\([^()]*\)\s*$", cur)
                if not m:
                    raise SystemExit("ui250: bracketed sentence not found")
                val = cur[:m.start()] + ver
            else:
                val = ver
            setv(key, val, f"translator's version ({SRC}, sheet 1)", "translator's version (whole cell)"
                 if key != "ui250" else "translator's version (the bracketed sentence)",
                 f"sheet 1 {key}: {ans}" + (f" — comment: {d['comment']}" if d["comment"] else ""))
        elif ans.startswith("Keep"):
            reg.add("handoff", key, "kept (translator: keep my wording)", cur, cur, f"sheet 1 {key}: {ans}")
        else:
            raise SystemExit(f"sheet 1 {key}: unrecognised answer {ans!r}")

    # ---- sheet 2 ---------------------------------------------------------
    for key, d in s2.items():
        ans = d["answer"]
        if key.startswith("ui.") :
            continue                       # frames: handled on the framework
        if key not in rows:
            raise SystemExit(f"sheet 2: {key} not in the handoff")
        cur = rows[key][5].value or ""
        if ans.startswith("Use the suggested"):
            val = cur
            for old, new in SHEET2_SUBSTITUTIONS[key]:
                val = sub_once(val, old, new, f"sheet 2 {key}")
            setv(key, val, f"translator chose the suggested wording ({SRC}, sheet 2)",
                 "suggested wording → applied as the exact substitution", f"sheet 2 {key}: {ans}")
        elif ans.startswith("Use my version"):
            ver = typo(d["version"], reg.typo, f"sheet 2 {key}")
            setv(key, ver, f"translator's version ({SRC}, sheet 2)", "translator's version (whole cell)",
                 f"sheet 2 {key}: {ans}")
        elif ans.startswith("Keep"):
            reg.add("handoff", key, "kept (translator: keep as it is)", cur, cur,
                    f"sheet 2 {key}: {ans}" + (f" — {d['comment']}" if d["comment"] else ""))
        else:
            raise SystemExit(f"sheet 2 {key}: unrecognised answer {ans!r}")

    # ---- sheet 3: page text that had no Welsh ------------------------------
    for key, d in s3.items():
        if key in FRAME_KEYS_FROM_SHEET3:
            continue
        cy = d["cy"]
        if key in NEW_HANDOFF_ROWS:
            cat, en = NEW_HANDOFF_ROWS[key]
            n = ws.max_row + 1
            ws.cell(row=n, column=1).value = n - 1
            ws.cell(row=n, column=2).value = cat
            ws.cell(row=n, column=3).value = "Translator"
            ws.cell(row=n, column=4).value = key
            ws.cell(row=n, column=5).value = en
            ws.cell(row=n, column=6).value = typo(cy, reg.typo, f"sheet 3 {key}")
            ws.cell(row=n, column=7).value = (f"row added V4.15: the payload declared this key with no "
                                             f"Strings row to land in ({SRC}, sheet 3)")
            rows[key] = tuple(ws.cell(row=n, column=c) for c in range(1, 8))
            reg.add("handoff", key, "NEW ROW (key declared by the payload, no row before)", None,
                    ws.cell(row=n, column=6).value, f"sheet 3 {key}")
            continue
        if key not in rows:
            raise SystemExit(f"sheet 3: {key} not in the handoff")
        cur = rows[key][5].value
        if cur:
            raise SystemExit(f"sheet 3 {key}: the handoff already carries Welsh ({cur[:40]!r})")
        if key == "ui054":
            # the translator gave the three example words in quotation marks;
            # the row's own note fixed the "·" separators — assembled, flagged
            words = re.findall(r'"([^"]+)"', cy)
            if len(words) != 3:
                raise SystemExit(f"ui054: expected three quoted words, got {words}")
            val = " · ".join(words)
            reg.flag("ui054", "example banner assembled from the translator's three quoted words",
                     f"cell: {cy!r} → {val!r}. The live chip for a selected sport reads "
                     f"'Dewiswyd Pêl Droed' (sheet 58 prefix 'Dewiswyd {{x}}'); the example says 'Wedi dewis Pêl droed'.",
                     "confirm the example wording, or align it with the live chip")
        elif key == "ui055":
            # AW-4 (item 8): "{n} ymateb disgybl wedi’u cynnwys" is the banner
            # count everywhere; the cell wrote the singular participle after 7
            val = "7 ymateb disgybl wedi’u cynnwys"
            reg.flag("ui055", "banner example: AW-4 form applied over the cell's participle",
                     f"cell: {cy!r}; AW-4 (item 8) chose '{{n}} ymateb disgybl wedi’u cynnwys' for the banner count — "
                     f"applied with n = 7: {val!r}", "confirm")
        else:
            val = typo(cy, reg.typo, f"sheet 3 {key}")
        setv(key, val, f"translated ({SRC}, sheet 3)", "NEW Welsh (row was untranslated)", f"sheet 3 {key}")

    # ---- AW-4: one term for each concept, where page text carries it -------
    for key, r in rows.items():
        v = r[5].value
        if not isinstance(v, str):
            continue
        if has_ltc(v):
            new = ltc_sub(v)
            setv(key, new, f"AW-4 term: {LTC_OLD} → {LTC_NEW} ({SRC})", f"term: {LTC_OLD} → {LTC_NEW}", "sheet 2 AW-4 item 1")
        if "y wedd hon" in (r[5].value or ""):
            v2 = r[5].value
            new = v2.replace("y wedd hon", "y golwg hwn")
            setv(key, new, f"AW-4 term: y wedd hon → y golwg hwn ({SRC})", "term: y wedd hon → y golwg hwn",
                 "sheet 2 AW-4 item 7 (‘Y golwg hwn’ for ‘this view’ everywhere)")
    # the remaining 'gwedd' (a view / any view) in the translator's prose is left verbatim — flagged below

    # ---- change-log sheet ---------------------------------------------------
    if "V4.15 change log" in hwb.sheetnames:
        del hwb["V4.15 change log"]
    wl = hwb.create_sheet("V4.15 change log")
    wl.append([f"Handoff V4.15 — generated {TODAY} by pipeline/apply_translator_handover.py from V4.14 and the {SRC}."])
    wl.append(["Key", "Change", "Before", "After", "Source"])
    for a in reg.applied:
        if a["target"] == "handoff":
            wl.append([a["key"], a["change"], a["before"], a["after"], a["source"]])
    return rows


# ================================================================ FRAMEWORK
def apply_framework(fwb, s2, s2b, s3, s4, s5, aw, reg: Register):
    log = []          # (sheet, key, change, before, after, reason)

    def L(sheet, key, change, before, after, reason):
        log.append([sheet, key, change, before, after, reason])
        reg.add(f"framework {sheet}", key, change, before, after, reason)

    # ---- 43 Interface frames ----------------------------------------------
    ws = fwb["43 Interface frames"]
    frames = {str(r[0].value): r for r in ws.iter_rows(min_row=4) if r[0].value}
    for key in FRAME_KEYS_FROM_SHEET3:
        d = s3[key]
        r = frames[key]
        cy = typo(d["cy"], reg.typo, f"sheet 3 {key}")
        en = (r[3].value or "").strip()
        if cy.strip() == en.strip():
            # the translator returned the English unchanged — not a translation
            reg.flag(key, "translator returned the English text unchanged",
                     f"sheet 3 row {key}: Welsh cell = {cy!r} (identical to the English). Left PENDING "
                     f"(shown in marked English when the accessibility switch is on).",
                     "re-ask the translator for the Welsh heading")
            r[7].value = f"PENDING — translator returned the English unchanged ({TODAY}); re-ask"
            L("43 Interface frames", key, "PENDING (English returned)", r[4].value, None, f"sheet 3 {key}")
            continue
        if r[4].value and (r[7].value or "").upper().startswith("EXCEPTION"):
            before = f"(EXCEPTION {r[7].value}: English served as data)"
        else:
            before = r[4].value
        r[4].value = cy
        old_status = r[7].value
        r[7].value = (f"TRANSLATED {TODAY} (handover sheet 3)"
                      + ("; Q17 exception closed" if (old_status or "").upper().startswith("EXCEPTION") else ""))
        L("43 Interface frames", key, "EDIT (translated)", before, cy,
          f"sheet 3 {key}" + (f" — comment: {d['comment']}" if d.get("comment") else ""))
    for key, subs in SHEET2_FRAME_SUBSTITUTIONS.items():
        d = s2[key]
        r = frames[key]
        if d["answer"].startswith("Use the suggested") and subs:
            val = r[4].value
            for old, new in subs:
                val = sub_once(val, old, new, f"sheet 2 {key}")
            L("43 Interface frames", key, "EDIT (suggested wording chosen)", r[4].value, val, f"sheet 2 {key}: {d['answer']}")
            r[4].value = val
            r[7].value = f"PROPOSED — edited {TODAY} (handover sheet 2)"
        else:
            L("43 Interface frames", key, "CONFIRMED (no edit)", r[4].value, r[4].value, f"sheet 2 {key}: {d['answer']}")

    # ---- 23 Answer labels -------------------------------------------------
    ws = fwb["23 Answer labels"]
    labels = {str(r[0].value): r for r in ws.iter_rows(min_row=4) if r[0].value}
    by_n = {d["n"]: d for d in s2b}

    def set_label(en, cy, reason, note=None):
        r = labels[en]
        before = r[1].value
        r[1].value = cy
        r[2].value = cy
        if note:
            r[8].value = ((r[8].value or "") + " " + note).strip()
        L("23 Answer labels", en, "EDIT", before, cy, reason)

    # 2b rows 43–45: "Not at all" on the confidence grid — the survey's own
    # option for that question (sheet 30 p.24) via the D32 grid mechanism
    set_label("Not at all", "Dim o gwbl  |  Ddim yn hyderus o gwbl  [confidence grid]",
              "sheet 2b rows 43–45 (Use survey wording): the confidence question's survey option is "
              "‘Ddim yn hyderus o gwbl’ (sheet 30, Q24); enjoyment / PE-feel keep ‘Dim o gwbl’",
              note=f"v2.7 ({TODAY}): D32 branches — default ‘Dim o gwbl’ (enjoyment, PE and Active Lessons; survey p.23, 26), "
                   "confidence grid ‘Ddim yn hyderus o gwbl’ (survey p.24). Translator: use survey wording (handover 2b rows 43–45).")
    reg.flag("Not at all (confidence chips)", "sheet 2b showed the wrong survey question",
             "rows 43–45 column I showed ‘Dim o gwbl’, which is the ENJOYMENT question's option; the confidence "
             "question the chips belong to (survey Q24) offers ‘Ddim yn hyderus o gwbl’. ‘Use survey wording’ was "
             "applied with the confidence question's own option.",
             "confirm ‘Hyder i roi cynnig ar gamp newydd: Ddim yn hyderus o gwbl’ (and the two sibling chips / charts)")
    # 2b row 12: "Other" on the take-part chip = the survey option (‘Arall, rho fanylion’)
    set_label("Other", "Arall  |  Arall, rho fanylion  [take-part grid]",
              "sheet 2b row 12 (Use survey wording): take-part question option ‘Arall, rho fanylion’ (sheet 30, Q6a); "
              "every other ‘Other’ keeps ‘Arall’ (row 23: keep today's)",
              note=f"v2.7 ({TODAY}): D32 branches — take-part grid ‘Arall, rho fanylion’ (survey p.6a), default ‘Arall’.")
    reg.flag("Other (take-part chip)", "English label is ‘Other’; Welsh now the full survey option",
             "row 12 chose the survey wording ‘Arall, rho fanylion’ for ‘Takes part: Other’; the English report label "
             "abbreviates the survey's ‘Other, please specify.’ to ‘Other’. Applied as decided.",
             "confirm the asymmetry is intended")
    # new rows: the ethnicity profile table (long labels) and "I don’t know"
    def add_label(en, cy, cat, gender, initial, mutable, status, note):
        n = ws.max_row + 1
        vals = [en, cy, cy, cat, gender, initial, mutable, status, note]
        for c, v in enumerate(vals, 1):
            ws.cell(row=n, column=c).value = v
        L("23 Answer labels", en, "NEW ROW", None, cy, note)

    add_label("I don’t know", "Dydw i ddim yn gwybod", "Scale", "—", "d", "YES", "ATTESTED-SURVEY",
              f"v2.7 ({TODAY}): approved survey (S0) p.19, 20 — the frequency grid's option. Replaces the code alias "
              "‘Ddim yn gwybod’ (welsh_render LABEL_ALIASES, never seen by a translator). Handover 2b row 58 / AW-1: survey wording.")
    r49 = by_n[49]
    add_label("White Welsh, English, Scottish, Northern Irish, or British", typo(r49["ans_tr"]), "Ethnicity", "—", "g", "YES",
              "TRANSLATOR", f"v2.7 ({TODAY}): profile table label. Translator's document 15 Sep 2026 (‘Your School’ table), chosen on "
              "handover 2b row 49 (Use document wording).")
    add_label("Asian, Asian Welsh, or Asian British", "Asiaidd, Asiaidd Cymreig, neu Asiaidd Prydeinig", "Ethnicity", "—", "vowel",
              "NO — immutable initial", "ATTESTED-SURVEY",
              f"v2.7 ({TODAY}): profile table label. Approved survey (S0) p.4; handover 2b row 51 (Use survey wording).")
    add_label("Black, Black Welsh, Black British, Caribbean, or African", "Du, Du Cymreig, Du Prydeinig, Caribïaidd, neu Affricanaidd",
              "Ethnicity", "—", "d", "YES", "ATTESTED-SURVEY",
              f"v2.7 ({TODAY}): profile table label. Approved survey (S0) p.4 — NOT on handover 2b (the row is never a chart "
              "option); applied by extension of rows 50–53 (survey wording for every survey-attested profile row). FLAGGED.")
    reg.flag("Black, Black Welsh, Black British, Caribbean, or African", "profile-table row not on sheet 2b",
             "the ethnicity profile table lists this survey option but the handover did not ask about it; the survey's Welsh "
             "‘Du, Du Cymreig, Du Prydeinig, Caribïaidd, neu Affricanaidd’ was applied by extension of the sibling decisions "
             "(rows 50–53: survey wording). The translator's document wrote it without the comma before ‘neu’.",
             "confirm")
    add_label("Other ethnic groups (for example, Arab)", "Grwpiau ethnig eraill [er enghraifft Arabaidd]", "Ethnicity", "g (m)", "g", "YES",
              "ATTESTED-SURVEY", f"v2.7 ({TODAY}): profile table label. Approved survey (S0) p.4, verbatim including the square "
              "brackets; handover 2b row 52 (Use survey wording).")
    reg.flag("Other ethnic groups (for example, Arab)", "square brackets from the survey",
             "row 52 chose the survey wording, which carries square brackets (‘[er enghraifft Arabaidd]’) where the English "
             "report label uses round brackets. Applied verbatim.", "confirm, or choose the round-bracket form")
    # AW-2 / AW-5 / AW-1: no label edits (survey wording kept) — recorded
    L("23 Answer labels", "*", "CONFIRMED", "survey wording", "survey wording",
      f"AW-1 ‘{aw['AW-1']['answer']}’; AW-2 ‘{aw['AW-2']['answer']}’; AW-5 ‘{aw['AW-5']['answer']}’")

    # ---- 58 Chip prefixes --------------------------------------------------
    ws = fwb["58 Chip prefixes"]
    chips = {(str(r[0].value), str(r[1].value)): r for r in ws.iter_rows(min_row=3) if r[0].value and r[1].value}

    def set_chip(kind, key, cy, status, reason):
        r = chips[(kind, key)]
        before = r[3].value
        if before == cy:
            r[4].value = status
            L("58 Chip prefixes", f"{kind} {key}", "CONFIRMED (same wording)", before, cy, reason)
            return
        r[3].value = cy
        r[4].value = status
        L("58 Chip prefixes", f"{kind} {key}", "EDIT", before, cy, reason)

    decided = {}
    for d in s2b:
        fam = PREFIX_EN_TO_FAM.get((d["pre_en"] or "").strip())
        if not fam or not d["dec_pre"] or d["dec_pre"] in ("—", "-"):
            continue
        if d["dec_pre"].startswith("Use document"):
            target = typo(d["pre_tr"])
        elif d["dec_pre"].startswith("Keep"):
            target = d["pre_today"]
        else:
            raise SystemExit(f"2b row {d['n']}: unrecognised prefix decision {d['dec_pre']!r}")
        if fam in decided and decided[fam][0] != target:
            raise SystemExit(f"2b: conflicting prefix decisions for {fam}: {decided[fam]} vs row {d['n']}")
        decided.setdefault(fam, (target, d["n"], d["dec_pre"]))
    today_prefix = {PREFIX_EN_TO_FAM[(d["pre_en"] or "").strip()]: d["pre_today"]
                    for d in s2b if (d["pre_en"] or "").strip() in PREFIX_EN_TO_FAM}
    for fam, (target, n, dec) in decided.items():
        cur = chips[("cohort prefix", fam)][3].value
        if not cur.startswith(today_prefix[fam]):
            raise SystemExit(f"sheet 58 {fam}: {cur!r} does not start with the handover's today-prefix {today_prefix[fam]!r}")
        suffix = cur[len(today_prefix[fam]):]          # ': {x}' or ' {x}'
        if fam == "cb" and target == decided.get("ov", (None,))[0]:
            # NOT APPLIED: the document prefix is the same phrase as the
            # ‘through sport’ chip's, so the two filter families would carry
            # one Welsh label for two English ones — the blocking gate
            # DIS-scope rejected the build (504 scope records). Today's
            # prefix (with ‘clwb’) is kept until the translator gives a
            # distinct one. Flagged.
            set_chip("cohort prefix", fam, cur, f"HELD {TODAY} — 2b row {n} ‘{dec}’ NOT applied: DIS-scope (see sheet 61)",
                     f"sheet 2b row {n}: {dec} — not applied (DIS-scope)")
            continue
        if fam == "ed":
            # row 4: answer ‘Oes’ + keep the prefix — but AW-3's comment says
            # keep the named form (PR-11), as do rows 47–48: the chip stays named
            reg.flag("Ethnically diverse background chip", "row 4 contradicts AW-3 (comment) and rows 47–48",
                     "2b row 4 chose ‘Use document wording’ for the answer (‘Oes’); AW-3's dropdown says ‘Use my document's "
                     "wording’ but its comment says ‘Keep the named form.’, and rows 47–48 keep ‘Cefndir ethnig amrywiol / "
                     "Heb fod o gefndir ethnig amrywiol’. The named form (PR-11) is kept; the chip's prefix is unchanged as decided.",
                     "confirm the named form for the filter group")
            set_chip("cohort prefix", fam, cur, f"TRANSLATOR {TODAY} (2b row {n}: keep)", f"sheet 2b row {n}: {dec}")
            continue
        new = target + suffix
        set_chip("cohort prefix", fam, new, f"TRANSLATOR {TODAY} (2b row {n}: {dec})", f"sheet 2b row {n}: {dec}")
    # cb: the document prefix omits ‘clwb’ — the same phrase as ov — NOT applied (DIS-scope)
    reg.flag("cb chip prefix (Estimated times active through club sport each week)",
             "NOT APPLIED — the translator's prefix is identical to the ‘through sport’ chip's",
             "2b row 8 chose ‘Use document wording’: ‘Amcan o’r amser y mae’n actif drwy chwaraeon bob wythnos’, which is "
             "exactly the prefix chosen for the ‘through sport’ chip (rows 6–7). Two English filters (‘…through sport each "
             "week: 3’ / ‘…through club sport each week: 3’) would then carry one Welsh label; the acceptance gate DIS-scope "
             "rejected that build (504 view labels). Today's prefix ‘Amcangyfrif o weithiau’n actif drwy chwaraeon clwb bob "
             "wythnos’ is kept for this chip only.",
             "give a distinct Welsh prefix for ‘through CLUB sport’ (the ‘through sport’ chip already has the document's)")
    # mi: ‘os’ + ‘Pe bai …’ — as decided
    reg.flag("mi chip prefix (Would do more sport if)", "prefix ends in ‘os’ before a ‘Pe bai …’ answer",
             "2b rows 17–23 chose the document prefix ‘Byddai’n gwneud mwy o chwaraeon os’ and the survey answers "
             "(‘Pe bai mwy o chwaraeon rydw i’n eu hoffi’ …): the chip reads ‘… os: Pe bai …’. Applied as decided.",
             "confirm")
    # ec and cn: not on sheet 2b — the translator's document form applied by
    # extension of the sibling decisions (ep/eb/eo, ct/cs/cg all → document)
    set_chip("cohort prefix", "ec", "Yn mwynhau clybiau’r tu allan i’r ysgol: {x}",
             f"TRANSLATOR document (by extension of 2b rows 24–38) — CONFIRM", "not on sheet 2b; document row ‘Enjoys clubs outside of school: …’")
    reg.flag("ec chip prefix (Enjoys clubs outside school)", "not asked on sheet 2b — applied by extension",
             "the three sibling enjoyment prefixes (PE, school clubs, other settings) were all switched to the translator's "
             "document form; the document's form for this one, ‘Yn mwynhau clybiau’r tu allan i’r ysgol’, was applied for consistency.",
             "confirm")
    set_chip("cohort prefix", "cn", "Hyder i roi cynnig ar gamp mewn lle newydd: {x}",
             f"TRANSLATOR document (by extension of 2b rows 43–45) — CONFIRM", "not on sheet 2b; document row ‘Confidence trying sports in a new place: …’")
    reg.flag("cn chip prefix (Confidence trying sport in a new place)", "not asked on sheet 2b — applied by extension",
             "the three sibling confidence prefixes were confirmed as the document's form (identical to today's); the document's "
             "form for this one differs (‘ar gamp mewn lle newydd’ for ‘ar chwaraeon mewn lle newydd’) and was applied for consistency.",
             "confirm")
    # gy: the long-term-condition term
    for fam in ("dy", "gy"):
        r = chips[("cohort prefix", fam)]
        if has_ltc(r[3].value):
            set_chip("cohort prefix", fam, ltc_sub(r[3].value), f"TRANSLATOR {TODAY} (AW-4: {LTC_NEW})",
                     "sheet 2 AW-4 item 1")
    # gender / banner labels
    r2 = by_n[2]
    set_chip("gender label", "all", typo(r2["ans_tr"]), f"TRANSLATOR {TODAY} (2b row 2: {r2['dec_ans']})",
             f"sheet 2b row 2 (All Respondents): {r2['dec_ans']}")
    r1 = by_n[1]
    set_chip("scope label", "whole", r1["ans_today"], f"TRANSLATOR {TODAY} (2b row 1: {r1['dec_ans']})",
             f"sheet 2b row 1 (Whole School): {r1['dec_ans']}")
    set_chip("banner label", "all", "Pob disgybl", f"kept (2b row 1/2: the banner's ‘All pupils’ is not the filter's ‘All Respondents’)",
             "banner audience part — unchanged")

    # ---- 31 Yes-No forms (dy form label) ------------------------------------
    ws = fwb["31 Yes-No forms"]
    for r in ws.iter_rows(min_row=2):
        for c in r:
            if has_ltc(c.value) and not c.value.startswith("Oes gen ti"):
                before = c.value
                c.value = ltc_sub(c.value)
                L("31 Yes-No forms", str(r[0].value), "EDIT (term)", before, c.value, "sheet 2 AW-4 item 1; 2b row 3 (prefix: document)")

    # ---- 22 / 44: the term in the qualifier and singleton frames -----------
    for sheet in ("22 Qualifier frames", "44 Singleton frames"):
        ws = fwb[sheet]
        for r in ws.iter_rows(min_row=2):
            for c in r:
                if has_ltc(c.value):
                    before = c.value
                    c.value = ltc_sub(c.value)
                    L(sheet, f"{r[0].value} / {(r[1].value or '')[:50]}", "EDIT (term)", before, c.value, "sheet 2 AW-4 item 1")

    # ---- 17 Core nouns: the term as a row ----------------------------------
    ws = fwb["17 Lexicon core nouns"]
    n = ws.max_row + 1
    for c, v in enumerate(["long-term condition", LTC_NEW, "g (m)", "cyflyrau hirdymor", "c", "Yes",
                           f"v2.7 ({TODAY}): the translator's term (AW-4 item 1; sheet 2 disability_condition). Supersedes D33 "
                           "(‘cyflwr tymor hir’, the survey's wording) — D86. Read by welsh.term() for the e5 sentences."], 1):
        ws.cell(row=n, column=c).value = v
    L("17 Lexicon core nouns", "long-term condition", "NEW ROW", None, LTC_NEW, "sheet 2 AW-4 item 1 / disability_condition")

    # ---- 11 Conjunctions: CONJ-06 ---------------------------------------------
    ws = fwb["11 Conjunctions"]
    r15 = next(x for x in s4 if x["n"] == 15)
    r12 = next(x for x in s5 if x["n"] == "12")
    for r in ws.iter_rows(min_row=4):
        if r[0].value == "CONJ-06":
            before = r[2].value
            r[2].value = ("MODE=figure — a figure written in digits is a SYMBOL: it takes ‘a’ (never ‘ac’) and never mutates; "
                          "words keep CONJ-01..04. (Translator's ruling, handover sheet 4 rule 15: readers read the figures as "
                          "decimal numbers, ‘ac 36’ is wrong; sheet 5 #12: ‘… 12 eu bod yn hyderus a 11 eu bod yn barod i ddysgu’.) "
                          "MODE=decimal (a/ac by the decimal reading: ac 11, a 36, ac 8) and MODE=vigesimal (the v1.5–v2.6 rule) "
                          "remain available to the engine; only the marker changes.")
            r[3].value = "a"
            r[4].value = "a 11; a 36; a 8; rhwng Blynyddoedd 3 a 11 — but ‘ac un ar ddeg’ when written in words"
            r[5].value = (f"CHANGED v2.7 ({TODAY}) — was: read the figure as its vigesimal word (ac 21, ac 11). The rule's "
                          f"own example on the handover was ‘ac 36’; the translator marked it WRONG: ‘{r15['comment']}’. "
                          f"Their corrected sentence #12 writes ‘a 11’ — which is not the decimal reading either (‘un deg un’ "
                          f"starts with a vowel), so the operative rule that satisfies both answers is the FIGURE rule; flagged for confirmation.")
            L("11 Conjunctions", "CONJ-06", "RULE CHANGED (MODE=figure)", before, r[2].value,
              f"sheet 4 rule 15: {r15['verdict']} — {r15['comment']}; sheet 5 #12: {r12['verdict']} — {r12['corrected']}")
    reg.flag("CONJ-06 (a/ac before a figure)", "rule comment and corrected sentence differ",
             "sheet 4 rule 15: WRONG, ‘readers will be reading the numbers as decimal numbers … the rule should reflect this’ "
             "(so ‘ac 36’ → ‘a 36’). Sheet 5 #12 corrects ‘ac 11’ → ‘a 11’ — but the decimal reading of 11 (‘un deg un’) begins "
             "with a vowel and would take ‘ac’. The only rule that produces both of the translator's outputs is ‘a before every "
             "figure written in digits’ (figures as symbols, as the translator also reasons for mutation in rule 17: ‘written as a "
             "figure, there would be no mutation’). Implemented as MODE=figure; ‘y’ before figures was already the engine's practice.",
             "confirm ‘a’ before every figure (a 11, a 8, a 36), or specify the decimal-reading rule (ac 11, ac 8, a 36)")

    # ---- 09 Ordinals: ail + soft mutation ------------------------------------
    ws = fwb["09 Ordinals"]
    r16 = next(x for x in s4 if x["n"] == 16)
    r34 = next(x for x in s5 if x["n"] == "34")
    for r in ws.iter_rows(min_row=4):
        if str(r[0].value) == "2":
            before = r[4].value
            r[4].value = ((before or "") + f" · v2.7 ({TODAY}): ‘ail’ takes the SOFT MUTATION of the following noun in both "
                          f"genders — yr ail gamp, yr ail leoliad, yr ail ateb. Translator: sheet 4 rule 16 WRONG "
                          f"(‘{r16['comment']}’), sheet 5 #34 ‘{r34['corrected']}’. Engine: welsh_render.ail_np().").strip(" ·")
            L("09 Ordinals", "2 (ail)", "RULE NOTE (mutation after ail)", before, r[4].value,
              f"sheet 4 rule 16: {r16['verdict']}; sheet 5 #34: {r34['verdict']}")

    # ---- 39 Provisional rulings: confirmations ----------------------------------
    ws = fwb["39 Provisional rulings"]
    confirmed = {"PR-01": 35, "PR-02": 10, "PR-03": 11, "PR-04": 21, "PR-05": 27, "PR-06": 36, "PR-07": 37,
                 "PR-08": 4, "PR-10": 28, "PR-11": 40, "PR-13": 21, "PR-14": 11}
    verdicts = {x["n"]: x["verdict"] for x in s4}
    for r in ws.iter_rows(min_row=4):
        pid = str(r[0].value)
        if pid in confirmed and verdicts.get(confirmed[pid]) == "Correct":
            before = r[5].value
            r[5].value = f"{before}; confirmed by the translator {TODAY} (handover sheet 4 rule {confirmed[pid]}: Correct)"
            L("39 Provisional rulings", pid, "CONFIRMED", before, r[5].value, f"sheet 4 rule {confirmed[pid]}")

    # ---- 02 Project decisions ----------------------------------------------
    ws = fwb["02 Project decisions"]
    decisions = [
        ("D84", "Numbers", "How is a figure written in digits read for a / ac?",
         "vigesimal reading (v1.5–v2.6: ac 11, ac 36, a 12); decimal reading (ac 11, a 36); figure as a symbol (a before every figure)",
         "FIGURE AS A SYMBOL: ‘a’ before every figure written in digits (a 11, a 36, a 8); no mutation of or after a figure; words keep CONJ-01..04 (MODE=figure on sheet 11 CONJ-06).",
         "The translator ruled CONJ-06 wrong (‘ac 36’) and corrected ‘ac 11’ to ‘a 11’; the figure rule is the one rule that yields both. Flagged for confirmation against the decimal-reading alternative.",
         f"{SRC}: sheet 4 rule 15, sheet 5 #12", "Translator (rule); report owner (application)"),
        ("D85", "Ordinals", "Does ‘ail’ mutate the noun that follows it?",
         "no mutation in the runner-up templates (‘Yr ail camp’, 2,465 sentences); soft mutation (‘Yr ail gamp’)",
         "SOFT MUTATION after ‘ail’ in both genders: yr ail gamp, yr ail leoliad (welsh_render.ail_np; _ord_np already did this).",
         "Translator: sheet 4 rule 16 WRONG — ‘Camp should mutate after ail’; sheet 5 #34 ‘Yr ail gamp mwyaf cyffredin oedd …’.",
         f"{SRC}: sheet 4 rule 16, sheet 5 #34", "Translator"),
        ("D86", "Terminology", "One term for ‘long-term condition’ everywhere.",
         "cyflwr tymor hir (the approved survey, D33); cyflwr hirdymor (the translator's document)",
         "cyflwr hirdymor — everywhere the report writes it (chip prefixes, chart heading, qualifier and singleton frames, the e5 sentences, sheet 31 form label). The survey text on sheet 30 is quoted unchanged.",
         "AW-4 item 1 ‘Cyflwr hirdymor’; sheet 2 disability_condition ‘Use the suggested wording’. Supersedes D33 for the report's own text.",
         f"{SRC}: sheet 2 AW-4, disability_condition", "Translator"),
        ("D87", "Answer labels and chips", "Whose wording do answer labels and chip prefixes carry?",
         "the survey's own option wording; the translator's document; today's report wording — per row (handover sheet 2b)",
         "ANSWER LABELS: the survey's wording (AW-1, AW-2, AW-5, 2b column O) — including the profile table's long labels (sheet 23 rows added v2.7) and the confidence grid's ‘Ddim yn hyderus o gwbl’ (D32 branch). CHIP PREFIXES: the translator's document where chosen (2b column P), otherwise today's; sheet 58 is the record. The filter option ‘All Respondents’ is ‘Pob Ymatebwr’; the banner's ‘All pupils’ stays ‘Pob disgybl’.",
         "Every row decided by the translator on sheet 2b; contradictions with AW-4 (‘Ddim yn siŵr’, ‘gennyf’, ‘Sefyll / Eistedd’) resolved for the per-row survey decision and flagged (sheet 61).",
         f"{SRC}: sheet 2b, AW-1..5", "Translator"),
    ]
    for d in decisions:
        ws.append(list(d))
        L("02 Project decisions", d[0], "NEW DECISION", None, d[4], d[6])

    # ---- 60 change log, 61 flags ---------------------------------------------
    for name in ("60 v2.7 change log", "61 V4.15 flags"):
        if name in fwb.sheetnames:
            del fwb[name]
    ws = fwb.create_sheet("60 v2.7 change log")
    ws.append([f"Framework v2.7 change log — generated {TODAY} by pipeline/apply_translator_handover.py from v2.6 and the {SRC}."])
    ws.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    for row in log:
        ws.append(row)
    ws = fwb.create_sheet("61 V4.15 flags")
    ws.append([f"Flags raised while applying the {SRC} — each is a question for the translator or the owner; the report shows what the "
               "‘action’ column's first clause describes. Generated by pipeline/apply_translator_handover.py."])
    ws.append(["Where", "What", "Detail", "Action needed"])
    for f in reg.flags:
        ws.append([f["key"], f["what"], f["detail"], f["action"]])
    ws0 = fwb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = (
        f"Version 2.7 ({TODAY}): the translator's completed handover applied — sheet 43 frames translated (sheet 3), sheet 23 labels "
        "(profile table rows, ‘I don’t know’, the confidence grid's ‘Ddim yn hyderus o gwbl’, the take-part ‘Arall, rho fanylion’), "
        "sheet 58 chip prefixes (sheet 2b), ‘cyflwr hirdymor’ (sheets 17/22/31/44/58), CONJ-06 MODE=figure (sheet 11), ‘ail’ + soft "
        "mutation (sheet 09), decisions D84–D87, twelve rulings confirmed (sheet 39); sheet 60 change log; sheet 61 flags.")
    return log


# ================================================================== register
def write_register(path, reg: Register, s4, s5, aw, s2b):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.append([f"SSS2026 — V4.15: the translator's completed handover applied ({TODAY})"])
    ws.append(["Every change made to the translation handoff (V4.15) and the Framework (v2.7) from the translator's answers, with its source row; "
               "then every flag: a contradiction between two answers, a typo left word for word, a decision applied by extension of its siblings, "
               "and the engine-rule rulings. Generated by pipeline/apply_translator_handover.py — nothing here was typed by hand."])
    ws.append([])
    ws.append(["Sheet", "Rows"])
    ws.append(["Applied", len(reg.applied)])
    ws.append(["Flags", len(reg.flags)])
    ws.append(["Typography", len(reg.typo)])
    ws.append(["Rule verdicts", len(s4) + len(s5)])
    hdr_font = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for title, rows, cols in (
        ("Applied", [[a["target"], a["key"], a["change"], a["before"], a["after"], a["source"]] for a in reg.applied],
         ["Target", "Key", "Change", "Before", "After", "Source"]),
        ("Flags", [[f["key"], f["what"], f["detail"], f["action"]] for f in reg.flags],
         ["Where", "What", "Detail", "Action needed"]),
        ("Typography", [[t["where"], t["kind"], t["before"], t["after"]] for t in reg.typo],
         ["Where", "Kind", "Before", "After"]),
        ("Rule verdicts", [[f"sheet 4 rule {x['n']}", x["area"], x["verdict"], x["comment"] or "", ""] for x in s4]
                          + [[f"sheet 5 #{x['n']}", "sentence", x["verdict"], x["why"] or "", x["corrected"] or ""] for x in s5],
         ["Row", "Area", "Verdict", "Comment", "Corrected sentence"]),
    ):
        w = wb.create_sheet(title)
        w.append(cols)
        for c in w[1]:
            c.font = hdr_font
            c.fill = fill
        for r in rows:
            w.append(r)
        for col, width in zip("ABCDEF", (28, 34, 60, 60, 60, 60)):
            w.column_dimensions[col].width = width
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def doc_white_label(docx_path):
    """The first Welsh line of the translator's 'Your School' ethnicity
    table (the White long form) — read from the document itself."""
    import docx
    d = docx.Document(docx_path)
    for t in d.tables:
        for r in t.rows:
            c = r.cells
            if len(c) >= 3 and c[1].text.strip().startswith("White Welsh, English, Scottish"):
                return c[2].text.strip().split("\n")[0].strip()
    raise SystemExit("translator document: the ethnicity table row was not found")


def main():
    completed, fw_in, fw_out, ho_in, ho_out, reg_out, docx_path = sys.argv[1:8]
    cwb = openpyxl.load_workbook(completed, data_only=True)
    s1 = read_sheet1(cwb)
    s2, aw = read_sheet2(cwb)
    s2b = read_sheet2b(cwb)
    s3 = read_sheet3(cwb)
    s4 = read_sheet4(cwb)
    s5 = read_sheet5(cwb)
    white = doc_white_label(docx_path)
    for d in s2b:
        if d["n"] == 49:
            if not white.startswith(typo(d["ans_tr"])):
                raise SystemExit(f"2b row 49: the handover cell {d['ans_tr']!r} is not a prefix of the document's line {white!r}")
            d["ans_tr"] = white          # the handover quoted the document truncated
    reg = Register()
    # ---- flags that come from reading, not from applying ----------------------
    reg.flag("AW-4 ‘Ddim yn siŵr’", "AW-4 vs sheet 2b rows 11 / 53 and AW-1",
             "AW-4 chose ‘Ddim yn siŵr’ as the one form; rows 11 and 53 chose the survey wording ‘Dydw i ddim yn siŵr’ for the "
             "answer option ‘I’m not sure’ and AW-1 keeps survey wording for quoted options. Applied: ‘Not sure’ = ‘Ddim yn siŵr’ "
             "(the survey's option), ‘I’m not sure’ = ‘Dydw i ddim yn siŵr’ (the survey's option) — the per-row decisions.",
             "confirm the two survey forms stay distinct")
    reg.flag("AW-4 ‘Mae’n well gennyf beidio â dweud’", "AW-4 vs sheet 2b row 46 and AW-1",
             "AW-4 chose ‘gennyf’; row 46 chose the survey wording ‘Mae’n well gen i beidio â dweud’ for the answer option and "
             "AW-1 keeps survey wording. Applied: the answer label keeps ‘gen i’; the translator's own page text keeps ‘gennyf’ verbatim.",
             "confirm")
    reg.flag("AW-4 ‘Sefyll / Eistedd’", "AW-4 vs sheet 2b rows 9, 10, 59, 60 and AW-1",
             "AW-4 chose ‘Sefyll / Eistedd’; the four per-row decisions and AW-1 keep the survey's ‘Yn sefyll / Yn eistedd’. "
             "Applied: the survey wording (labels and chips unchanged).", "confirm")
    reg.flag("AW-4 ‘Y golwg hwn’ vs the translator's ‘gwedd’", "term applied only to the exact phrase",
             "‘y wedd hon’ → ‘y golwg hwn’ applied in ui193 (‘Argraffu neu arbed y golwg hwn’) and ui264. The translator's other "
             "uses of ‘gwedd’ (ui034 ‘gwedd newydd ar y data’, ui197 ‘Yng ngwedd yr ysgol gyfan’, ui255 ‘gwedd ddemograffig’, "
             "‘y wedd bresennol’, ui264 ‘unrhyw wedd’) are their prose and are left word for word.",
             "say whether those should also become ‘golwg’")
    reg.flag("ui086", "typos left word for word", "the translator's version carries ‘pennodol’, ‘gwestynnau’, ‘I’r’ (capital I) and "
             "‘a ddefnyddiwyd y lleoliad yno’; applied exactly as written (only the straight apostrophe was curled).",
             "confirm or correct")
    reg.flag("ui189", "typo left word for word", "‘Diolch a cammau nesaf’ — ‘cammau’ (camau / chamau?) applied as written.", "confirm or correct")
    reg.flag("ui230", "typos left word for word", "‘disgybyl’, ‘yn ol’, ‘amylder’ applied as written.", "confirm or correct")
    reg.flag("ws_combined", "double space left word for word", "‘Yn siarad Cymraeg  (grŵp hidlo cyfun)’ — two spaces before the bracket (collapsed by the browser).", "confirm")
    reg.flag("dl_combined / settings_dl", "two forms of ‘and/or’", "the filter group says ‘neu/ac anhawster dysgu’ (dl_combined) and the chart heading ‘a/neu anawsterau dysgu’ (settings_dl); both applied as written.", "confirm")
    reg.flag("ui.gdpr_body_module", "‘for the current view’ = ‘i’r farn bresennol’", "applied as written; elsewhere ‘this view’ is ‘y golwg hwn’ (AW-4).", "confirm")
    reg.flag("ui.chart_sports_total_*", "‘dywedon nhw’ / ‘dywedasant’", "the three sibling notes use ‘y dywedon nhw eu bod’ (club) and ‘y dywedasant eu bod’ (community, other); applied as written.", "confirm the register")
    reg.flag("ui255 vs AW-4", "‘gwedd’ in the translator's version", "the ‘Use my version’ text keeps ‘gwedd ddemograffig’ and ‘y wedd bresennol’; AW-4 chose ‘golwg’ for ‘this view’. Left as written (see the ‘gwedd’ flag).", "confirm")

    hwb = openpyxl.load_workbook(ho_in)
    apply_handoff(hwb, s1, s2, s3, aw, reg)
    hwb.save(ho_out)
    fwb = openpyxl.load_workbook(fw_in)
    apply_framework(fwb, s2, s2b, s3, s4, s5, aw, reg)
    fwb.save(fw_out)
    write_register(reg_out, reg, s4, s5, aw, s2b)
    Path(reg_out).with_suffix(".json").write_text(
        json.dumps({"applied": reg.applied, "flags": reg.flags, "typography": reg.typo}, ensure_ascii=False, indent=1),
        encoding="utf-8")
    print(f"handoff → {ho_out}; framework → {fw_out}; register → {reg_out}")
    print(f"applied {len(reg.applied)} changes · {len(reg.flags)} flags · {len(reg.typo)} typographic normalisations")


if __name__ == "__main__":
    main()

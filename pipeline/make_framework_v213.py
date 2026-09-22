# -*- coding: utf-8 -*-
"""Framework v2.13 — the workbook edit for V6.0 (production, chunked delivery).

    python -m pipeline.make_framework_v213 config/01_Framework_v2.12.xlsx config/01_Framework_v2.13.xlsx

Sheet 43 only: THREE new interface frames for the loading mask the served
package shows while a view's state chunks are fetched and verified —
``ui.loading_view`` (the status message), ``ui.load_error`` (the message when
a chunk cannot be loaded) and ``ui.retry`` (the retry button). English
written by Industryline (plain status text, nothing about findings); Welsh
EMPTY → TRANSLATOR (pending: shown as marked English in Welsh mode until
returned). The monolithic assurance copy never shows the mask. Change log
sheet 72; V6.0 flags sheet 73. No Welsh composed; no translator wording
changed; no narrative surface touched.
"""
from __future__ import annotations

import datetime
import sys

import openpyxl

FRAMES = [
    ("ui.loading_view", "#load-mask — status message while a view's data is fetched and verified (role=status)",
     "Loading this view…"),
    ("ui.load_error", "#load-mask — message when the view's data could not be loaded (fail-closed: nothing is shown)",
     "This view could not be loaded. Check your connection and try again."),
    ("ui.retry", "#load-retry — the button that retries the load", "Try again"),
]


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    assert "72 v2.13 change log" not in wb.sheetnames
    ws43 = wb["43 Interface frames"]
    keys = {str(r[0].value) for r in ws43.iter_rows(min_row=4) if r[0].value}
    for key, where, en in FRAMES:
        assert key not in keys, key
        assert en == en.strip() and en[0] not in "·—–:;," and en[-1] not in "·—–:;,"
        ws43.append([key, "V6.0 chunked delivery", where, en, None, "—",
                     "V6.0 (22 Sep 2026): the served report loads each view's data on demand from verified chunk "
                     "files; while that happens the content is masked and this text is shown. Plain status text; "
                     "no finding is stated. Welsh pending the translator.", "TRANSLATOR"])
    h = ws43.cell(row=1, column=1)
    h.value = str(h.value) + " · v2.13: +ui.loading_view, +ui.load_error, +ui.retry (Welsh pending) — see sheet 72."
    ws72 = wb.create_sheet("72 v2.13 change log")
    ws72.append([f"Framework v2.13 change log — generated {datetime.date.today().isoformat()} by pipeline/make_framework_v213.py from v2.12. "
                 "Sheet 43 only: the three loading-mask frames of the served (chunked) package. No Welsh composed; no translator wording changed; no narrative surface touched."])
    ws72.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    for key, _where, en in FRAMES:
        ws72.append(["43 Interface frames", key, "NEW ROW", "", en, "V6.0 — production chunked delivery (deployment plan §2, review P0.7)"])
    ws73 = wb.create_sheet("73 V6.0 flags")
    ws73.append(["Flags raised while preparing V6.0 — each is a question for the translator, Sport Wales or the owner; nothing here is decided by Industryline."])
    ws73.append(["Where", "What", "Detail", "Action needed"])
    ws73.append(["43 Interface frames — ui.loading_view, ui.load_error, ui.retry", "Welsh for the three loading-mask strings",
                 "New frames; English only. Shown as marked English in Welsh mode until returned. They appear only on the served package while a view loads.",
                 "Translator: three rows in the next handoff (sheet 3 'Translate'), with the two frames still pending (ui.a11y_switch_desc, ui.a11y_glossary_heading)."])
    ws73.append(["49 English client edits", "Owner confirmation of the three English strings",
                 "Plain status text written by Industryline: 'Loading this view…', 'This view could not be loaded. Check your connection and try again.', 'Try again'.",
                 "Owner: confirm or amend the wording (a change is a workbook edit and a rebuild)."])
    ws0 = wb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = ("Version 2.13 (V6.0, 22 Sep 2026): sheet 43 +ui.loading_view, +ui.load_error, +ui.retry (the served package's "
                                       "loading mask; Welsh pending the translator); sheet 72 change log; sheet 73 V6.0 flags. No translator wording changed.")
    wb.save(dst)
    print(f"written {dst}: {len(FRAMES)} sheet-43 rows added")


if __name__ == "__main__":
    main()

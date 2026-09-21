# -*- coding: utf-8 -*-
"""Framework v2.4 — the workbook edit for the V4.13 (accessibility) round.

    python -m pipeline.make_framework_v24 config/01_Framework_v2.3.xlsx config/01_Framework_v2.4.xlsx

Three new sheet-43 frames carry the ONLY new reader-facing text of V4.13:
the accessibility switch's label and the heading and one-line note of the
glossary panel the switch reveals (the glossary's entries are the guide's
existing FAQ questions — no new definitions). English only, status
TRANSLATOR: no Welsh word is composed here; the client shows the English
under the review marker (lang="en") in Welsh mode until the translator's
handoff returns the Welsh, exactly as the V4.11 pending frames do.

Everything else the accessibility switch does is presentation (fonts,
sizes, colours, opened data tables, captions taken from the existing
ui.alt_* frames) and needs no text. Logged on sheet "56 v2.4 change log".
"""
from __future__ import annotations

import datetime
import sys

import openpyxl

PEND = "TRANSLATOR"

# (key, path, consumer, en, cy, slots, notes, status)
NEW_FRAMES = [
    ("ui.a11y_switch", "V4.13 accessibility", "#btn-a11y switch label (rail, below the language slider)",
     "Accessibility features", "", "—",
     "Visible label and accessible name of the accessibility switch. Off by default; on, the report uses Aptos/Arial, larger text, higher-contrast colours, a colour-blind-safe chart palette, opened data tables, captions under the character images and the glossary panel — in both languages, content unchanged.", PEND),
    ("ui.a11y_glossary_heading", "V4.13 accessibility", "#a11y-glossary heading (shown only when the switch is on)",
     "Glossary of terms used in this report", "", "—",
     "Feedback 'Plain English and layout': a glossary of technical terms. The entries are links to the guide's existing FAQ answers (base, percentages, colours, grey bars, hidden numbers, Free School Meal quartiles, the banner) — no new definitions.", PEND),
    ("ui.a11y_glossary_intro", "V4.13 accessibility", "#a11y-glossary note",
     "Select a term to open its explanation in the guide.", "", "—", "", PEND),
]


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    ws = wb["43 Interface frames"]
    existing = {str(r[0].value) for r in ws.iter_rows(min_row=4) if r[0].value}
    log = []
    for f in NEW_FRAMES:
        assert f[0] not in existing, f[0]
        ws.append(list(f))
        log.append(("43 Interface frames", f[0], "NEW", "", "(empty — translator)", f[6]))
    ws56 = wb.create_sheet("56 v2.4 change log")
    ws56.append([f"Framework v2.4 change log — generated {datetime.date.today().isoformat()} by pipeline/make_framework_v24.py from v2.3. Source: 'School Reports Accessibility Feedback.docx' (sha 71dffd0e…, 17 Sep 2026)."])
    ws56.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    for row in log:
        ws56.append(list(row))
    ws0 = wb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = ("Version 2.4 (V4.13 accessibility round, 17 Sep 2026): three sheet-43 frames for the accessibility switch and its glossary panel "
                                      "(English, awaiting the translator); sheet 56 change log. Nothing else changed.")
    wb.save(dst)
    print(f"written {dst}: {len(NEW_FRAMES)} new frames")


if __name__ == "__main__":
    main()

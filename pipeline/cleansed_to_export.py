# -*- coding: utf-8 -*-
"""One school's rows of the cleansed 2026 pupil dataset, written in the
SmartSurvey "Responses wide" layout the report pipeline reads (V5.0
real-school round, pipeline 0.28.0).

    python -m pipeline.cleansed_to_export <stage2_full_cleaned.parquet> <school_id> <out.xlsx>
                                          [--s16 config/S16_Translations_MASTER_v6_English_import.xlsx]

Why this exists. The V4.15 prototype was built from a SmartSurvey export of
one school (columns "Q_<question id>_<text>", "; "-joined multi-selects,
JSON-fragment matrix cells). The cleansed Stage 2 v2 dataset (Industryline
cleansing handover v2.3, 9 Sep 2026) holds the same answers for every school
in a flat, coded layout. This module is the inverse of that flattening for
one school: it writes each row back in the prototype's layout so the report
pipeline — loader, derivations, engine, narrative, Welsh generation, gates —
runs UNCHANGED on real schools. It changes no answer: every value is a
codebook label copied through, every sport label is the live survey's own
option text (S16), and every derived measure is left to the pipeline.

What is deliberately NOT copied: free text (never leaves the dataset — an
"Other" answer becomes the bare "Other" prefix the loader already recodes),
IP address and entry URL (the prototype excluded them; the dataset has none).

Completion status: the dataset's analytical set is "completed OR partial
reaching page 48+" (the data owner's inclusion rule). Rows are written with
Completion_Status "Complete" / "Partial" as recorded; a school profile lists
the statuses it accepts (acceptedStatuses), so the loader's rule stays
explicit and logged.

Reproducibility: rows are written in response_id order; the Metadata sheet
records the source file's sha256, the school filter and the row counts.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import html as _html
import json
import re
import sys
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq

from .common import CONFIG_DIR

# ---- the prototype's question ids -------------------------------------------
Q = {
    "school_confirm": 26871527, "gender": 26871497, "age": 26871555, "year": 26871554,
    "ethnicity": 26871498, "eth_white": 26871499, "eth_mixed": 26871500, "eth_asian": 26871501,
    "eth_black": 26871502, "eth_other": 26871503,
    "condition_1": 26871504, "condition_2": 26871505, "welsh": 26871567, "welsh_sport": 26871590,
    "demand": 26871551, "enjoy": 26871552, "confidence": 26871586, "pe_feel": 26871587,
    "join_in": 26871565, "listened": 26871566, "more_if": 26871588, "important": 26871589,
    "take_part": 26871506, "wheelchair": 26871557,
}
# category order: fitness, team, water, cycling, games/target, performance,
# racket/paddle, adventure, combat, snow/ice
CATEGORY_KEYS = ["fitness", "team", "water", "cycling", "games", "performance",
                 "racket", "adventure", "combat", "snow"]
CAT_MAIN_Q = dict(zip(CATEGORY_KEYS, range(26871507, 26871517)))     # the ten tick questions
CAT_OTHER_Q = {"fitness": 26871578, "team": 26871553, "water": 26871579, "games": 26871580,
               "performance": 26871581, "racket": 26871582, "adventure": 26871583,
               "combat": 26871584, "snow": 26871585}                    # the "other sports" pages
CAT_SETTING_Q = dict(zip(CATEGORY_KEYS, range(26871517, 26871527)))
FREQ_Q = {
    "pe":             dict(zip(CATEGORY_KEYS, range(26871531, 26871541))),
    "school_club":    dict(zip(CATEGORY_KEYS, range(26871591, 26871601))),
    "outside_club":   dict(zip(CATEGORY_KEYS, range(26871541, 26871551))),
    "somewhere_else": dict(zip(CATEGORY_KEYS, range(26871568, 26871578))),
}
SETTING_TITLE = {"pe": "In PE or Lesson Time", "school_club": "In a school club",
                 "outside_club": "In a club outside of school", "somewhere_else": "Somewhere else"}
# S16 (the live instrument, the Vale of Glamorgan master import) question ids
# of the same questions — the option lists are the same master
S16_MAIN = {"fitness": "26872035", "team": "26872036", "water": "26872037", "cycling": "26872038",
            "games": "26872039", "performance": "26872040", "racket": "26872041",
            "adventure": "26872042", "combat": "26872043", "snow": "26872044"}
S16_OTHER = {"fitness": "26872106", "team": "26872081", "water": "26872107", "games": "26872108",
             "performance": "26872109", "racket": "26872110", "adventure": "26872111",
             "combat": "26872112", "snow": "26872113"}
S16_DEMAND = "26872079"
NON_SPORT = {"None of these", "My sport isn’t here"}

# the codebook's label strings that differ typographically from the
# prototype export's (straight vs curly apostrophe) — mapped by CODE, never
# by string surgery
FREQ_BY_CODE = {1: "Three or more times a week", 2: "2 times a week", 3: "1 time a week",
                4: "Less than once a week", 5: "I don’t know"}
HOWMUCH_BY_CODE = {1: "A lot", 2: "A little", 3: "Not much", 4: "Not at all", 5: "Not sure"}
SCALE_BY_CODE = {1: "Very", 2: "Quite", 3: "Not very", 4: "Not at all"}
ENJOY_ROWS = [("pe_lessons", "PE lessons"), ("school_sports_clubs", "School sports clubs"),
              ("sports_clubs_outside_of_school", "Sports clubs outside of school"),
              ("other_settings_like_in_the_park_or_garden", "Other settings, like in the park or garden")]
CONF_ROWS = [("try_a_new_sport", "Try a new sport?"), ("learn_a_new_skill", "Learn a new skill?"),
             ("try_again_when_sport_is_hard", "Try again when sport is hard?"),
             ("try_sport_in_a_new_place", "Try sport in a new place?")]
FEEL_ROWS = [("healthy", "Healthy?"), ("confident", "Confident?"), ("ready_to_learn", "Ready to learn?")]
MORE_IF = [("there_were_more_sports_i_liked", "There were more sports I liked"),
           ("i_felt_more_confident", "I felt more confident"),
           ("it_was_easier_to_take_part", "It was easier to take part"),
           ("i_felt_more_motivated", "I felt more motivated"),
           ("it_felt_more_comfortable_for_me", "It felt more comfortable for me"),
           ("i_had_what_i_need_or_it_cost_less", "I had what I need or it cost less"),
           ("none_of_these", "None of these")]
IMPORTANT = [("having_fun", "Having fun"), ("being_with_friends", "Being with friends"),
             ("improving_my_skills", "Improving my skills"), ("feeling_confident", "Feeling confident"),
             ("winning", "Winning"), ("liking_my_coach_or_teacher", "Liking my coach or teacher"),
             ("feeling_safe", "Feeling safe"),
             ("access_to_good_equipment_and_facilities", "Access to good equipment and facilities"),
             ("none_of_these", "None of these")]
TAKE_PART = [("standing", "Standing"), ("seated", "Seated"),
             ("with_communication_aids", "With communication aids"), ("none_of_these", "None of these"),
             ("i_m_not_sure", "I’m not sure"), ("prefer_not_to_say", "Prefer not to say")]
WHEELCHAIR = [("on_the_floor_or_on_a_chair", "On the floor or on a chair"),
              ("i_use_a_manual_wheelchair_for_sport_and_at_school", "I use a manual wheelchair for sport and at school"),
              ("i_only_use_a_wheelchair_for_sport", "I only use a wheelchair for sport"),
              ("using_a_power_chair", "Using a power chair"), ("none_of_these", "None of these")]
# The codebook codes a detail answer given through the "Any other …" option's
# text box as "Other (free text)" (the last code of each detail question);
# the raw export carried it as "Any other … background / <text>". Verified on
# the prototype school: its 11 "Any other White background / …" answers are
# exactly the dataset's 11 "Other (free text)" rows. The category is restored
# (the owner rule counts the CATEGORY, never the text); the text is withheld.
ANY_OTHER = {"ethnicity_white_detail": "Any other White background",
             "ethnicity_mixed_detail": "Any other Mixed or Multiple background",
             "ethnicity_asian_detail": "Any other Asian background",
             "ethnicity_black_detail": "Any other Black, Black Welsh, Black British, Caribbean or African background",
             "ethnicity_other_detail": "Any other ethnic group"}


def clean(s):
    s = re.sub(r"<[^>]+>", " ", str(s))
    s = _html.unescape(s)
    return re.sub(r"\s+", " ", s).strip()


def cslug(label):
    """The cleansing pipeline's column slug for an option label."""
    s = label.lower().replace("’", "'")
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def s16_structure(path):
    """{category: {"main": [labels], "other": [labels]}}, demand labels, and
    slug -> label for every sport option — from the live instrument."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Translations"].iter_rows(values_only=True))
    opts = {}
    for r in rows[3:]:
        m = re.match(r"o_(\d+)_(\d+)_text$", str(r[0] or ""))
        if m and r[2] is not None:
            opts.setdefault(m.group(1), []).append(clean(r[2]))
    cats = {}
    for k in CATEGORY_KEYS:
        cats[k] = {"main": [o for o in opts[S16_MAIN[k]] if o not in NON_SPORT],
                   "other": [o for o in opts.get(S16_OTHER.get(k, ""), []) if o not in NON_SPORT]}
    demand = [o for o in opts[S16_DEMAND]]
    return cats, demand


def matrix_cell(pairs):
    """The SmartSurvey matrix-cell encoding the loader parses (row_title /
    column_title JSON fragments joined by '; ')."""
    if not pairs:
        return None
    return "; ".join(json.dumps({"type": "matrix_row", "row_title": r, "column_title": c},
                                ensure_ascii=False) for r, c in pairs)


def _repack_deterministic(path, pinned):
    """openpyxl stamps every zip entry with the wall clock; rewrite the
    container with the pinned time and a fixed entry order so the file's
    bytes depend on its content only."""
    import zipfile
    src = zipfile.ZipFile(path, "r")
    entries = sorted(src.namelist())
    data = {n: src.read(n) for n in entries}
    src.close()
    # openpyxl re-stamps dcterms:modified at save time — pin it too
    core = data["docProps/core.xml"].decode("utf-8")
    core = re.sub(r'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)',
                  lambda m: m.group(1) + pinned.strftime("%Y-%m-%dT%H:%M:%SZ") + m.group(2), core)
    data["docProps/core.xml"] = core.encode("utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        for n in entries:
            zi = zipfile.ZipInfo(n, date_time=pinned.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, data[n])


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("parquet"); ap.add_argument("school_id"); ap.add_argument("out")
    ap.add_argument("--s16", default=str(CONFIG_DIR / "S16_Translations_MASTER_v6_English_import.xlsx"))
    a = ap.parse_args(argv)
    cats, demand_labels = s16_structure(a.s16)
    pf = pq.ParquetFile(a.parquet)
    names = pf.schema_arrow.names
    # sport slug -> (label, category, tier); the cleansing slugs are the
    # option labels slugged, a few truncated — matched exactly, else by a
    # unique prefix, never guessed
    sport_by_slug = {}
    for k, d in cats.items():
        for tier, labels in (("main", d["main"]), ("other", d["other"])):
            for lab in labels:
                sport_by_slug[cslug(lab)] = (lab, k, tier)
    done_slugs = [n[len("sports_done__"):] for n in names if n.startswith("sports_done__")]
    slug_map = {}
    for s in done_slugs:
        if s in ("none_of_these", "my_sport_isn_t_here"):
            continue
        if s in sport_by_slug:
            slug_map[s] = sport_by_slug[s]
        else:
            cand = [k for k in sport_by_slug if k.startswith(s)]
            if len(cand) != 1:
                raise SystemExit(f"no unique S16 option for dataset sport column {s!r}: {cand}")
            slug_map[s] = sport_by_slug[cand[0]]
    demand_by_slug = {cslug(l): l for l in demand_labels}
    try_slugs = [n[len("sports_try__"):] for n in names if n.startswith("sports_try__")]
    for s in try_slugs:
        if s not in demand_by_slug:
            raise SystemExit(f"no S16 demand option for dataset column sports_try__{s}")
    setting_cols = {(st, s): f"sport_setting_{st}__{s}" for st in FREQ_Q for s in slug_map
                    if f"sport_setting_{st}__{s}" in names}
    freq_cols = {(st, s): f"sport_freq_{st}__{s}_code" for st in FREQ_Q for s in slug_map
                 if f"sport_freq_{st}__{s}_code" in names}
    # ---- read the school's rows -------------------------------------------
    tbl = pq.read_table(a.parquet, filters=[("school_id", "=", a.school_id)])
    if tbl.num_rows == 0:
        raise SystemExit(f"no rows for school_id {a.school_id}")
    tbl = tbl.sort_by("response_id")
    rows = tbl.to_pylist()
    src_sha = hashlib.sha256(Path(a.parquet).read_bytes()).hexdigest()
    headers = json.loads((CONFIG_DIR / "export_headers.json").read_text(encoding="utf-8"))
    qhead = {q["qid"]: q["header"] for q in headers["questions"]}
    admin = headers["admin"]

    def v(r, k):
        x = r.get(k)
        return None if x is None or (isinstance(x, float) and x != x) else x

    def ts(x):
        if x is None:
            return None
        if isinstance(x, dt.datetime):
            return x.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        return str(x)

    def multi(r, prefix, items, other_text_col=None):
        vals = [v(r, f"{prefix}__{k}") for k, _ in items]
        if all(x is None or x == -1 for x in vals):
            return None
        chosen = [lab for (k, lab), x in zip(items, vals) if x == 1]
        if other_text_col and v(r, other_text_col):
            chosen.append("Other")
        return "; ".join(chosen) if chosen else None

    out_rows = []
    n_by_status = {}
    for r in rows:
        rec = {}
        status = "Complete" if v(r, "status") == "completed" else "Partial"
        n_by_status[status] = n_by_status.get(status, 0) + 1
        rec["Response_ID"] = v(r, "response_id")
        rec["Completion_Status"] = status
        rec["School_Name"] = v(r, "school_name"); rec["School_ID"] = v(r, "school_id")
        rec["Local_Authority"] = v(r, "local_authority"); rec["Survey_ID"] = v(r, "survey_id")
        rec["Tracking_Link_ID"] = v(r, "tracking_link_id")
        rec["Date_Started"] = ts(v(r, "date_started")); rec["Date_Modified"] = ts(v(r, "date_modified"))
        rec["Date_Ended"] = ts(v(r, "date_ended")); rec["Response_Date"] = ts(v(r, "response_date"))
        rec["Current_Page_ID"] = v(r, "current_page_id"); rec["Language"] = v(r, "language_translation_id")
        q = {}
        q[Q["school_confirm"]] = v(r, "school_confirm"); q[Q["gender"]] = v(r, "gender")
        q[Q["age"]] = v(r, "age"); q[Q["year"]] = v(r, "year_group_clean")
        q[Q["ethnicity"]] = v(r, "ethnicity")
        for k, col in (("eth_white", "ethnicity_white_detail"), ("eth_mixed", "ethnicity_mixed_detail"),
                       ("eth_asian", "ethnicity_asian_detail"), ("eth_black", "ethnicity_black_detail"),
                       ("eth_other", "ethnicity_other_detail")):
            x = v(r, col)
            q[Q[k]] = (ANY_OTHER[col] if x == "Other (free text)" else x) if x is not None else None
        q[Q["condition_1"]] = v(r, "condition_1"); q[Q["condition_2"]] = v(r, "condition_2")
        q[Q["welsh"]] = v(r, "welsh_ability"); q[Q["welsh_sport"]] = v(r, "yesno_2")
        q[Q["join_in"]] = v(r, "enjoy_pe"); q[Q["listened"]] = v(r, "take_part")
        # ---- sports: the ten tick questions and the "other sports" pages ----
        done = {s: v(r, f"sports_done__{s}") for s in slug_map}
        block_answered = any(x is not None and x != -1 for x in done.values()) or \
            v(r, "sports_done__none_of_these") == 1 or v(r, "sports_done__my_sport_isn_t_here") == 1
        for k in CATEGORY_KEYS:
            main = [slug_map[s][0] for s in slug_map if slug_map[s][1] == k and slug_map[s][2] == "main" and done[s] == 1]
            other = [slug_map[s][0] for s in slug_map if slug_map[s][1] == k and slug_map[s][2] == "other" and done[s] == 1]
            if block_answered:
                q[CAT_MAIN_Q[k]] = "; ".join(main) if main else "None of these"
                if k in CAT_OTHER_Q:
                    q[CAT_OTHER_Q[k]] = "; ".join(other) if other else None
            else:
                q[CAT_MAIN_Q[k]] = None
                if k in CAT_OTHER_Q:
                    q[CAT_OTHER_Q[k]] = None
        # ---- the settings grids and the four frequency grids ----------------
        for k in CATEGORY_KEYS:
            pairs = []
            for s, (lab, cat, _tier) in slug_map.items():
                if cat != k:
                    continue
                for st in FREQ_Q:
                    c = setting_cols.get((st, s))
                    if c and v(r, c) == 1:
                        pairs.append((lab, SETTING_TITLE[st]))
            q[CAT_SETTING_Q[k]] = matrix_cell(pairs)
            for st, qmap in FREQ_Q.items():
                fp = []
                for s, (lab, cat, _tier) in slug_map.items():
                    if cat != k:
                        continue
                    c = freq_cols.get((st, s))
                    code = v(r, c) if c else None
                    if code is not None:
                        fp.append((lab, FREQ_BY_CODE[int(code)]))
                q[qmap[k]] = matrix_cell(fp)
        # ---- demand ---------------------------------------------------------
        tv = {s: v(r, f"sports_try__{s}") for s in try_slugs}
        if all(x is None or x == -1 for x in tv.values()):
            q[Q["demand"]] = None
        else:
            q[Q["demand"]] = "; ".join(demand_by_slug[s] for s in try_slugs if tv[s] == 1) or None
        # ---- grids ----------------------------------------------------------
        q[Q["enjoy"]] = matrix_cell([(lab, HOWMUCH_BY_CODE[int(v(r, f"grid_howmuch__{k}_code"))])
                                     for k, lab in ENJOY_ROWS if v(r, f"grid_howmuch__{k}_code") is not None])
        q[Q["confidence"]] = matrix_cell([(lab, SCALE_BY_CODE[int(v(r, f"grid_confidence__{k}_code"))])
                                          for k, lab in CONF_ROWS if v(r, f"grid_confidence__{k}_code") is not None])
        q[Q["pe_feel"]] = matrix_cell([(lab, SCALE_BY_CODE[int(v(r, f"grid_feel__{k}_code"))])
                                       for k, lab in FEEL_ROWS if v(r, f"grid_feel__{k}_code") is not None])
        # ---- fixed multi-selects ---------------------------------------------
        q[Q["more_if"]] = multi(r, "why_more_active", MORE_IF, "why_more_active_other_text")
        q[Q["important"]] = multi(r, "what_matters", IMPORTANT, "what_matters_other_text")
        tp = multi(r, "mobility_position", TAKE_PART)
        tp_other = (v(r, "mobility_position__other_please_specify") == 1) or bool(v(r, "mobility_position_other_text"))
        mp_asked = any(v(r, f"mobility_position__{k}") not in (None, -1) for k, _ in TAKE_PART) or \
            v(r, "mobility_position__other_please_specify") not in (None, -1)
        if tp is None and tp_other:
            q[Q["take_part"]] = "Other (please specify):"
        elif tp is not None and tp_other:
            q[Q["take_part"]] = tp + "; Other (please specify):"
        else:
            q[Q["take_part"]] = tp if mp_asked else None
        q[Q["wheelchair"]] = multi(r, "wheelchair_use", WHEELCHAIR, "wheelchair_use_other_text")
        for qid in qhead:
            rec[qhead[qid]] = q.get(qid)
        out_rows.append(rec)
    # ---- write --------------------------------------------------------------
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Responses wide"
    cols = admin + [q["header"] for q in headers["questions"]]
    ws.append(cols)
    for rec in out_rows:
        ws.append([rec.get(c) for c in cols])
    wm = wb.create_sheet("Metadata")
    meta = [("source_file", Path(a.parquet).name), ("source_sha256", src_sha),
            ("filter", f"school_id == {a.school_id}"), ("rows_written", len(out_rows)),
            ("rows_by_status", json.dumps(n_by_status)),
            ("s16_file", Path(a.s16).name), ("s16_sha256", hashlib.sha256(Path(a.s16).read_bytes()).hexdigest()),
            ("written_by", "pipeline/cleansed_to_export.py (pipeline 0.28.0)"),
            ("note", "free text withheld; every answer is a codebook label copied through")]
    for k, val in meta:
        wm.append([k, val])
    out = Path(a.out); out.parent.mkdir(parents=True, exist_ok=True)
    # reproducible bytes: the workbook's own timestamps are pinned to the
    # cleansing handover's (v2.3, 9 Sep 2026 12:12:36), so the same rows
    # always give the same file and the same sourceChecksum in the build
    pinned = dt.datetime(2026, 9, 9, 12, 12, 36)
    wb.properties.created = pinned
    wb.properties.modified = pinned
    wb.properties.creator = "pipeline/cleansed_to_export.py"
    wb.properties.lastModifiedBy = "pipeline/cleansed_to_export.py"
    wb.save(out)
    _repack_deterministic(out, pinned)
    print(f"{out}: {len(out_rows)} rows ({n_by_status}); sports mapped {len(slug_map)}; demand options {len(try_slugs)}")


if __name__ == "__main__":
    main()

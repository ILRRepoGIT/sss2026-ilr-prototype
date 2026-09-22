# -*- coding: utf-8 -*-
"""Framework v2.9 — the workbook edit for V4.16: the Young Artists
Competition artwork replaces the Brain Break characters.

    python -m pipeline.make_framework_v29 config/01_Framework_v2.8.xlsx config/01_Framework_v2.9.xlsx

Nothing here composes Welsh, and no translator wording changes. The only
edits are on sheet 43 (interface frames):

  * fourteen new ``img alt`` frames, ``ui.alt_yac_*`` — one per artwork
    file, English alt text (a description of the picture; no pupil's name),
    Welsh EMPTY → status TRANSLATOR (PENDING: the client shows the marked
    English in Welsh mode until the translator returns the row);
  * the three Brain Break alt frames (``ui.alt_brain_*``) → DEPRECATED:
    the client no longer consumes them (FRAME-dormant reads the status).

Plus the change log (sheet 64) and the V4.16 flags for the translator
(sheet 65). The English of each frame is the single source the template's
alt literal and the client must match (tests/test_v416_images.py).
"""
from __future__ import annotations

import datetime
import sys

import openpyxl

from .yac_assets import ALT_EN, RETIRED, YAC

NOTE = ("V4.16 (22 Sep 2026): Young Artists Competition winning entry supplied by Sport Wales in the "
        "'Finals YCA' pack; alt text describes the picture only — no pupil's name, region or school. "
        "Welsh pending the translator (shown as marked English in Welsh mode until returned).")


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    assert "64 v2.9 change log" not in wb.sheetnames
    log = []
    ws43 = wb["43 Interface frames"]
    rows43 = {str(r[0].value): r for r in ws43.iter_rows(min_row=4) if r[0].value}
    for key in RETIRED:
        r = rows43[key]
        old_status = str(r[7].value or "")
        r[7].value = "DEPRECATED"
        r[6].value = ((r[6].value or "") + " DEPRECATED v2.9 (V4.16): the Brain Break characters were replaced by "
                      "the Young Artists Competition artwork; the client no longer consumes this frame.").strip()
        log.append(("43 Interface frames", key, "DEPRECATED", old_status, "DEPRECATED",
                    "V4.16 — Brain Break images replaced by the Young Artists Competition artwork (owner instruction, 22 Sep 2026)"))
    for key, fname, _token, _w, _crop, _bg, src_note in YAC:
        assert key not in rows43, key
        en = ALT_EN[key]
        assert en == en.strip() and en[0] not in "·—–:;," and en[-1] not in "·—–:;,", key
        ws43.append([key, "V4.16 artwork", f"img alt ({fname})", en, None, "—",
                     NOTE + " Alt text source: " + src_note + ".", "TRANSLATOR"])
        log.append(("43 Interface frames", key, "NEW ROW", "", en, "V4.16 — " + src_note))
    h = ws43.cell(row=1, column=1)
    h.value = str(h.value) + " · v2.9: +14 ui.alt_yac_* alt frames (Welsh pending), 3 ui.alt_brain_* DEPRECATED — see sheet 64."
    # ---- 64 change log ------------------------------------------------------
    ws64 = wb.create_sheet("64 v2.9 change log")
    ws64.append([f"Framework v2.9 change log — generated {datetime.date.today().isoformat()} by pipeline/make_framework_v29.py from v2.8. "
                 "Sheet 43 only: the Young Artists Competition artwork's alt frames (English written by Industryline from the artwork and, "
                 "for six entries, the survey's own mascot descriptions in S16; Welsh pending the translator) and the three Brain Break "
                 "frames deprecated. No Welsh composed; no translator wording changed; no narrative surface touched."])
    ws64.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    for row in log:
        ws64.append(list(row))
    # ---- 65 V4.16 flags ------------------------------------------------------
    ws65 = wb.create_sheet("65 V4.16 flags")
    ws65.append(["Flags raised while preparing V4.16 (the Young Artists Competition artwork) — each is a question for the translator or Sport Wales; nothing here is decided by Industryline."])
    ws65.append(["Where", "What", "Detail", "Action needed"])
    ws65.append(["43 Interface frames — ui.alt_yac_* (14 rows)", "Welsh alt text for the fourteen artwork images",
                 "The survey instrument (S16) carries the mascot descriptions in English only (the Welsh pages of the survey omit the mascot block), "
                 "so no attested Welsh exists for any of the fourteen alt texts. Until returned, the Welsh report shows the English alt text "
                 "marked as pending (lang=\"en\", 'Heb ei gyfieithu eto — dangosir y Saesneg'), as for every pending frame since V4.11.",
                 "Translator: fourteen rows in the next handoff (sheet 3 'Translate'). The alt text must stay a description of the picture — no pupil's name."])
    ws65.append(["43 Interface frames — ui.alt_brain_* (3 rows)", "Brain Break alt frames deprecated",
                 "The three character images were replaced by the artwork; their rows were never sent to the translator (withheld in the V4.14 handover because the images were being changed) and are now DEPRECATED.",
                 "None — recorded for completeness."])
    ws65.append(["Report — artwork publication", "Publication approval of the fourteen entries",
                 "The artwork is embedded as supplied. No visible pupil name or signature was observed in the files; the survey's own mascot pages named the artists, which the report does not.",
                 "Sport Wales: confirm publication approval and any credit line required before the artwork is used in a school-facing report."])
    ws65.append(["43 Interface frames — ui.a11y_glossary_heading", "Still pending (V4.15 flag 15)",
                 "The translator returned the English text unchanged; the row remains PENDING.",
                 "Translator: re-ask (carried forward from sheet 61)."])
    # ---- README -------------------------------------------------------------
    ws0 = wb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = ("Version 2.9 (V4.16 artwork, 22 Sep 2026): sheet 43 +14 ui.alt_yac_* alt frames for the Young Artists "
                                       "Competition entries (English from the artwork / S16 mascot descriptions; Welsh pending the translator), "
                                       "ui.alt_brain_* ×3 DEPRECATED; sheet 64 change log; sheet 65 V4.16 flags. No translator wording changed.")
    wb.save(dst)
    print(f"written {dst}: {len(YAC)} sheet-43 rows added, {len(RETIRED)} deprecated")


if __name__ == "__main__":
    main()

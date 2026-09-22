# -*- coding: utf-8 -*-
"""Framework v2.12 — the workbook edit for V4.19 (22 Sep 2026).

    python -m pipeline.make_framework_v212 config/01_Framework_v2.11.xlsx config/01_Framework_v2.12.xlsx

  sheet 43  the fourteen `ui.alt_yac_*` frames receive their Welsh, entered
            VERBATIM from the text the owner supplied in the session on
            22 Sep 2026 (status records the source; whether the translator
            produced it is for the owner to confirm — sheet 71). Nothing is
            composed here; the strings below are the owner's, character for
            character.
  sheet 49  EN-09 row: a correction appended (the "summary rows" it names
            never rendered; the chapter-summary sentence is the one citation).
  sheet 69  V4.18 flag row: the same correction appended.
  sheet 70  change log; sheet 71 V4.19 flags. No English changed.
"""
from __future__ import annotations

import datetime
import sys

import openpyxl

# The owner's text, verbatim (22 Sep 2026). Do not edit here: a change is a
# new workbook version with its own change-log row.
WELSH_ALT = {
    "ui.alt_yac_dragon_kit": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: draig goch yn gwisgo crys gwyn Cymru â’r rhif 11 arno, yn dal pêl rygbi, gydag un droed ar bêl droed a chenhinen yn ei chynffon",
    "ui.alt_yac_welsh_symbols": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: symbolau Cymru a chwaraeon gyda’i gilydd — cenhinen Bedr, pêl droed, bat a phêl griced, draig goch, pêl rygbi a chenhinen",
    "ui.alt_yac_balls": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: rhes o bedair pêl — pêl griced goch, pêl ddu â smotiau gwyn, pêl fasged oren a phêl ddu â’r rhif wyth arni",
    "ui.alt_yac_footballer": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: chwaraewr pêl droed â gwallt hir golau, mewn crys glas â’r rhif 7 arno, siorts gwyn ac esgidiau pinc, yn cicio pêl droed",
    "ui.alt_yac_football_splash": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: pêl droed â phaneli coch a gwyrdd yn ffrwydro allan o dasgiad o baent coch a gwyrdd",
    "ui.alt_yac_cyclist": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: beiciwr mynydd mewn helmed, yn plygu’n isel dros feic pinc wrth fynd dros naid, yn erbyn tasgiad glas o ddyfrlliw",
    "ui.alt_yac_dragon_football": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: draig goch sy’n gwenu, gydag adenydd lliw hufen, yn sefyll wrth ymyl pêl droed",
    "ui.alt_yac_heart": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: calon goch sy’n cynnwys pedair golygfa o ddraig goch yn cymryd rhan mewn pêl droed, rygbi, gymnasteg a chriced",
    "ui.alt_yac_horse": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: marchog mewn siaced goch a helmed ddu ar gefn ceffyl brown sy’n neidio",
    "ui.alt_yac_tennis_football": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: pêl droed uwchben pêl denis werdd lachar",
    "ui.alt_yac_gymnastics": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: clwb gymnasteg — un gymnastwr yn neidio uwchben trawst cydbwysedd, un arall yn hongian o far uchel, wrth ymyl pentwr o gylchoedd lliwgar",
    "ui.alt_yac_cricket": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: chwaraewr criced mewn helmed, menig a phadiau, yn siglo bat tuag at bêl goch",
    "ui.alt_yac_basketball": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: chwaraewr pêl fasged â chynffon ferlen, mewn fest goch a siorts du, yn neidio i daflu pêl fasged wedi’i fframio gan ffrwydrad coch siâp seren",
    "ui.alt_yac_dragon_wales": "Darlun o’r Gystadleuaeth Artistiaid Ifanc: draig goch Gymreig gyda’r gair WALES mewn llythrennau gwyrdd ar ei thraws",
}
STATUS = "TRANSLATED 2026-09-22 (supplied by the owner in the session; entered verbatim — translator provenance to confirm, sheet 71)"
CORRECTION = (" [CORRECTED v2.12 (22 Sep 2026): the summary cards ‘Club sport 3+ times a week (estimated)’ / "
              "‘Less than weekly club sport’ named here exist only in unreferenced generator code and have never "
              "rendered — the shipped V4.18 file carries neither string; the chapter-summary sentence "
              "‘… took part in club sport at least once a week’ is the only remaining citation of the club-sport "
              "estimate (EN-10 question unchanged).]")


def main():
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    assert "70 v2.12 change log" not in wb.sheetnames
    log = []
    ws43 = wb["43 Interface frames"]
    rows43 = {str(r[0].value): r for r in ws43.iter_rows(min_row=4) if r[0].value}
    for key, cy in WELSH_ALT.items():
        r = rows43[key]
        assert (r[4].value or "") == "" and str(r[7].value) == "TRANSLATOR", (key, r[4].value, r[7].value)
        r[4].value = cy
        r[7].value = STATUS
        r[6].value = ((r[6].value or "") + " Welsh entered v2.12 (22 Sep 2026) verbatim from the owner's text; "
                      "no pupil's name (description only).").strip()
        log.append(("43 Interface frames", key, "Welsh frame", "(empty — TRANSLATOR)", cy,
                    "V4.19 — owner-supplied Welsh alt text, 22 Sep 2026 (entered verbatim)"))
    h = ws43.cell(row=1, column=1)
    h.value = str(h.value) + " · v2.12: the fourteen ui.alt_yac_* frames carry their Welsh — see sheet 70."
    ws49 = wb["49 English client edits"]
    for r in ws49.iter_rows(min_row=1):
        if r[0].value == "EN-09":
            r[4].value = (r[4].value or "") + CORRECTION
            log.append(("49 English client edits", "EN-09", "correction appended (no change to the sanctioned edit)", "", CORRECTION.strip(), "V4.19"))
    ws69 = wb["69 V4.18 flags"]
    for r in ws69.iter_rows(min_row=3):
        if r[1].value and "club-sport estimate still appears" in str(r[1].value):
            r[2].value = (r[2].value or "") + CORRECTION
            log.append(("69 V4.18 flags", "row 1", "correction appended", "", CORRECTION.strip(), "V4.19"))
    ws70 = wb.create_sheet("70 v2.12 change log")
    ws70.append([f"Framework v2.12 change log — generated {datetime.date.today().isoformat()} by pipeline/make_framework_v212.py from v2.11. "
                 "Sheet 43: Welsh for the fourteen Young Artists Competition alt frames, entered verbatim from the owner's text of 22 Sep 2026. "
                 "Sheets 49/69: a correction appended. No English changed; nothing composed."])
    ws70.append(["Sheet", "Key", "Change", "Before", "After", "Reason / source"])
    for row in log:
        ws70.append(list(row))
    ws71 = wb.create_sheet("71 V4.19 flags")
    ws71.append(["Flags raised while preparing V4.19 — each is a question for the translator, Sport Wales or the owner; nothing here is decided by Industryline."])
    ws71.append(["Where", "What", "Detail", "Action needed"])
    ws71.append(["Sheet 43 — the fourteen ui.alt_yac_* frames", "Provenance of the Welsh",
                 "The Welsh was supplied by the owner in the session on 22 Sep 2026 and entered verbatim. The status row records that; it does not say who translated it.",
                 "Owner: confirm whether the translator produced it (then the status can read ‘TRANSLATED (translator, date)’) or whether the translator should review it."])
    ws71.append(["Sheet 43 — two frames still pending", "ui.a11y_switch_desc and ui.a11y_glossary_heading",
                 "Both still render as marked English in Welsh mode.", "Translator: return both (the glossary heading was returned as English unchanged on 21 Sep — re-ask)."])
    ws71.append(["Report — Summary page", "EN-10 (carried from V4.18)",
                 "The chapter-summary sentence ‘… took part in club sport at least once a week’ still cites the club-sport estimate.",
                 "Owner: confirm it stays, or sanction its removal."])
    ws0 = wb["00 README"]
    ws0.insert_rows(1)
    ws0.cell(row=1, column=1).value = ("Version 2.12 (V4.19, 22 Sep 2026): sheet 43 — Welsh for the fourteen Young Artists Competition alt frames, "
                                       "entered verbatim from the owner's text (provenance to confirm, sheet 71); sheets 49/69 — a correction appended "
                                       "(the ‘summary cards’ never rendered); sheet 70 change log; sheet 71 V4.19 flags. No English changed.")
    wb.save(dst)
    print(f"written {dst}")


if __name__ == "__main__":
    main()

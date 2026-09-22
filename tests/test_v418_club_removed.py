# -*- coding: utf-8 -*-
"""V4.18 (EN-09, owner instruction 22 Sep 2026): the Club Sports section is
gone from the report — module, chart, selected groups, appendix table, static
rows — and the metric behind it is still computed for the summary rows that
cite it. The behaviour of the weekly-frequency chart under a setting
selection (wider picture, definition sentence) is asserted on the built
package by the jsdom regression; here the structure and the sources."""
import json
import re
from pathlib import Path

import openpyxl
import yaml

from pipeline import narrative2_modules as NM
from pipeline import narrative2_modules_b as NB
from pipeline.common import latest_framework, CONFIG_DIR

ROOT = Path(__file__).resolve().parents[1]
TPL = (ROOT / "web" / "template.html").read_text(encoding="utf-8")
APP = (ROOT / "web" / "app.js").read_text(encoding="utf-8")
LEX = json.loads((ROOT / "config" / "welsh_lexicon.json").read_text(encoding="utf-8"))
METRICS = yaml.safe_load((ROOT / "config" / "metrics.yml").read_text(encoding="utf-8"))
RETIRED_ROWS = ("ui130", "ui131", "ui132", "ui134", "ui248")


def test_module_d2_is_no_longer_generated_or_labelled():
    assert "d2" not in NB.FullNarrator.MODULES
    assert "d2" not in NM.MODULE_LABELS
    assert "d0" in NB.FullNarrator.MODULES and "d3" in NB.FullNarrator.MODULES


def test_club_estimate_is_not_a_selected_group_source_but_stays_a_metric():
    cohort_sources = METRICS["cohorts"] if "cohorts" in METRICS else METRICS["cohort_sources"]
    assert "club_freq_estimate" not in cohort_sources
    assert "club_freq_estimate" in METRICS["metrics"]
    assert "participation_settings" in cohort_sources and "freq_estimate" in cohort_sources


def test_template_carries_no_club_sports_block():
    assert 'data-chart="club_freq_estimate"' not in TPL
    assert 'data-module="d2"' not in TPL
    assert "Club Sports" not in TPL
    assert "ui.appx_club" not in TPL and "ui.appx_club" not in APP


def test_appendix_club_frame_is_deprecated_not_deleted():
    assert "ui.appx_club" in LEX["deprecated_frames"]
    assert "ui.appx_club" not in LEX["interface_frames"]
    fw = latest_framework(CONFIG_DIR)
    assert fw.name == "01_Framework_v2.11.xlsx"
    wb = openpyxl.load_workbook(fw, read_only=True, data_only=True)
    rows = {str(r[0]): r for r in wb["43 Interface frames"].iter_rows(min_row=4, values_only=True) if r and r[0]}
    assert str(rows["ui.appx_club"][7]) == "DEPRECATED"
    en09 = [r for r in wb["49 English client edits"].iter_rows(values_only=True) if r and r[0] == "EN-09"]
    assert len(en09) == 1 and "SANCTIONED" in str(en09[0][5])


def test_five_static_rows_retired_with_the_translators_welsh_kept():
    keys = {r["key"] for r in LEX["handoff"]}
    for k in RETIRED_ROWS:
        assert k not in keys, k
    wb = openpyxl.load_workbook(CONFIG_DIR / "SSS2026_Welsh_Translation_Handoff_V4.18.xlsx", read_only=True)
    retired = {str(r[0]): r for r in wb["Retired"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    assert sorted(retired) == sorted(RETIRED_ROWS)
    prev = openpyxl.load_workbook(CONFIG_DIR / "SSS2026_Welsh_Translation_Handoff_V4.15.xlsx", read_only=True)["Strings"]
    prev_cy = {str(r[3]): (r[5] or "") for r in prev.iter_rows(min_row=2, values_only=True) if r and r[3]}
    for k, r in retired.items():
        assert "EN-09" in str(r[2])
        assert r[3] and r[3] == prev_cy[k], k      # the translator's Welsh, verbatim, on the record


def test_d0_leads_with_the_definition_sentence_under_a_setting_selection():
    src = (ROOT / "pipeline" / "narrative2_modules.py").read_text(encoding="utf-8")
    body = src[src.index("def m_d0("):src.index("def m_d2(")]
    assert re.search(r'if src_m == "participation_settings":\s*\n(\s*#[^\n]*\n)*\s*self\.para\(out, self\.group_note\(key, S, "d0"\)\)', body)
    # the sentence is the existing template, not a new one
    assert "group_defined_v2" in src


def test_client_keeps_the_wider_picture_for_the_frequency_chart_under_a_setting():
    assert re.search(r'participation_settings.*freq_estimate|freq_estimate.*participation_settings', APP, re.S)
    assert "wider" in APP and 'ui.ctx_showing' in APP
    assert 'const wholeShown = wider && state.scope === "whole" && state.gender === "all"' in APP

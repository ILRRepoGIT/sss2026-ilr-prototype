# -*- coding: utf-8 -*-
"""The recipient register: which schools receive a report, under which profile,
from the canonical dataset (review P0.2 "the definitive report universe";
deployment plan §4.2). One row per school in the dataset; the build starts
from this register and must finish with exactly the eligible rows built.

    python -m prod.register build <stage2.parquet> <out_dir> [--plasc <plasc2026.xlsx>] \\
                                  [--min-responses 5] [--framework config/01_Framework_v2.19.xlsx]
    python -m prod.register profile <register.json> <school_id> <out.json> --release <tag>

`build` writes register.json and register.csv (the same rows), a summary
(register_summary.json) and nothing else. Every value is either copied from
the dataset / PLASC or derived by a stated rule:

  name          PLASC 2026 school name when the school id matches, else the
                dataset's; name_cy = name (no Welsh school-name source is
                supplied — flagged for Sport Wales)
  la, la_cy     the dataset's local authority; Welsh from Framework sheet 53
                by exact English match (the dataset spellings without a
                sheet-53 row are recorded as la_cy_missing)
  partnership   from the dataset's region column: the five survey regions
                map one-to-one onto the five Regional Sport Partnerships
                named on sheet 53 (validated against the V5.1 profiles)
  years         the year groups with at least one accepted response, sorted;
                a gap inside the range is recorded (years_gap) and the report
                takes the EN-11 overview sentence (Framework v2.14)
  phase         primary if every year <= 6, secondary if every year >= 7,
                combined otherwise (the profile family of the prototype)
  eligible      n >= min_responses (default 1 — owner instruction 22 Sep 2026:
                every school with an accepted response receives a report; where
                the school has fewer than five the report's views are all
                suppressed by the rule of five inside the report itself, so the
                report is thin but exists and has its link) and no hold; the
                status column names the reason otherwise

Statuses: eligible · no_report_below_threshold · held_no_year ·
held_la_name_cy (any authority spelling without a sheet-53 row; since Framework
v2.15 carries the four dataset spellings — owner instruction 22 Sep 2026 — no
school is held on the 22 Sep dataset; --allow-missing-la-cy overrides). Special schools are built
under the profile family their years imply and flagged (special=true) for
Sport Wales's profile-matrix decision.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq

REGION_TO_PARTNERSHIP = {
    "North Wales": "Actif North Wales",
    "Mid Wales": "Mid Wales Sport Partnership [MWSP]",
    "South West Wales": "West Wales Sport Partnership [WWSP]",
    "South Wales Central": "Central South Active Partnership [CSAP]",
    "South East Wales": "Gwent Sport Partnership [GSP]",
}
FIELDWORK = "13th April – 17th July 2026"
PRIMARY_YEARS = [3, 4, 5, 6]
SECONDARY_YEARS = [7, 8, 9, 10, 11]


def slugify(name: str, school_id: str) -> str:
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    s = re.sub(r"-{2,}", "-", s)[:48].strip("-")
    return f"{s}-{school_id}" if s else f"school-{school_id}"


def proper_names(framework: Path) -> dict:
    wb = openpyxl.load_workbook(framework, read_only=True)
    ws = wb["53 Proper names"]
    out = {"Local authority": {}, "Regional Sport Partnership": {}}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r and r[0] in out and r[1]:
            out[r[0]][str(r[1]).strip()] = (str(r[2]).strip() if r[2] else "")
    return out


def plasc_names(plasc: Path) -> dict:
    """School ID -> (official name, sector, years with pupils on roll)."""
    wb = openpyxl.load_workbook(plasc, read_only=True)
    ws = wb["Final"]
    rows = ws.iter_rows(values_only=True)
    hdr = None
    out = {}
    for r in rows:
        if hdr is None:
            if r and r[0] == "Sector":
                hdr = list(r)
                idx = {h: i for i, h in enumerate(hdr) if h}
                ycols = {y: idx.get(f"All: Year {y}") for y in range(3, 12)}
            continue
        if not r or r[idx["School ID"]] in (None, ""):
            continue
        sid = str(r[idx["School ID"]]).strip()
        try:
            sid = str(int(float(sid)))
        except ValueError:
            pass
        years = []
        for y, ci in ycols.items():
            v = r[ci] if ci is not None else None
            try:
                if v is not None and str(v) not in ("[c]", "[z]", "*") and float(v) > 0:
                    years.append(y)
            except ValueError:
                pass
        out[sid] = {"name": str(r[idx["School"]]).strip(), "sector": str(r[idx["Sector"]]).strip(),
                    "la": str(r[idx["LA"]]).strip(), "years": years}
    return out


def scope_groups(years: list[int]) -> tuple[str, list[dict], str]:
    """Profile family, scope groups and the 'School Stages Covered' label."""
    p = [y for y in years if y <= 6]
    s = [y for y in years if y >= 7]
    yr = lambda ys: [{"key": f"y{y}", "label": f"Year {y}", "years": [y]} for y in ys]
    if p and s:
        fam = "combined"
        groups = ([{"key": "whole", "label": "Whole school", "description": f"Years {years[0]}–{years[-1]}", "years": years},
                   {"key": "primary", "label": "Primary phase", "description": f"Years {p[0]}–{p[-1]}", "years": p},
                   {"key": "secondary", "label": "Secondary phase", "description": f"Years {s[0]}–{s[-1]}", "years": s}]
                  + yr(years))
        stages = f"Primary and secondary (Years {years[0]}–{years[-1]})"
    elif p:
        fam = "primary"
        groups = [{"key": "whole", "label": "Whole school", "description": f"Years {p[0]}–{p[-1]}", "years": p}] + yr(p)
        stages = f"Primary (Years {p[0]}–{p[-1]})"
    else:
        fam = "secondary"
        groups = [{"key": "whole", "label": "Whole school", "description": f"Years {s[0]}–{s[-1]}", "years": s}] + yr(s)
        stages = f"Secondary (Years {s[0]}–{s[-1]})"
    if len(years) == 1:
        groups[0]["description"] = f"Year {years[0]}"
    return fam, groups, stages


def build(parquet: Path, out_dir: Path, plasc: Path | None, framework: Path, min_responses: int,
          hold_missing_la_cy: bool) -> dict:
    cols = ["response_id", "school_id", "school_name", "school_ref_status", "local_authority", "region",
            "status", "school_phase", "school_type", "year_group_num", "analytical_inclusion_flag"]
    t = pq.ParquetFile(parquet).read(columns=cols).to_pandas()
    t = t[t["analytical_inclusion_flag"].astype(bool)]
    names = proper_names(framework)
    pl = plasc_names(plasc) if plasc else {}
    rows = []
    no_school = t[t["school_ref_status"] != "Matched school"]
    t = t[t["school_ref_status"] == "Matched school"]
    for sid, g in t.groupby("school_id", sort=True):
        sid = str(sid)
        years = sorted({int(y) for y in g["year_group_num"].dropna()})
        by_year = {int(y): int(n) for y, n in g["year_group_num"].dropna().astype(int).value_counts().items()}
        by_status = {str(k): int(v) for k, v in g["status"].value_counts().items()}
        ds_name = str(g["school_name"].iloc[0]).strip()
        la = str(g["local_authority"].iloc[0]).strip()
        region = str(g["region"].iloc[0]).strip()
        p = pl.get(sid)
        name = p["name"] if p else ds_name
        la_cy = names["Local authority"].get(la, "")
        partnership = REGION_TO_PARTNERSHIP.get(region, "")
        gap = bool(years) and years != list(range(years[0], years[-1] + 1))
        n = int(len(g))
        status, reason = "eligible", ""
        if n < min_responses:
            status, reason = "no_report_below_threshold", f"{n} accepted responses; fewer than the --min-responses floor of {min_responses}"
        elif not years:
            status, reason = "held_no_year", "no accepted response carries a year group"
        elif hold_missing_la_cy and not la_cy:
            status, reason = "held_la_name_cy", f"Framework sheet 53 has no Welsh row for the local authority spelling '{la}'"
        fam, groups, stages = scope_groups(years) if years else ("primary", [], "")
        rows.append({
            "school_id": sid, "slug": slugify(name, sid), "name": name, "name_cy": name,
            "name_dataset": ds_name, "name_source": "PLASC 2026" if p else "dataset",
            "la": la, "la_cy": la_cy, "la_cy_missing": not la_cy, "region": region,
            "partnership": partnership, "partnership_cy": names["Regional Sport Partnership"].get(partnership, ""),
            "phase_dataset": str(g["school_phase"].iloc[0]), "type_dataset": str(g["school_type"].iloc[0]),
            "special": str(g["school_type"].iloc[0]) == "Special",
            "family": fam, "stages": stages, "years": years, "years_gap": gap,
            "plasc_years": (p or {}).get("years", []),
            "n": n, "n_by_status": by_status, "n_by_year": by_year,
            "status": status, "reason": reason, "scope_groups": groups,
        })
    summary = {
        "schools": len(rows), "responses": int(sum(r["n"] for r in rows)),
        "no_school_rows": int(len(no_school)),
        "no_school_by_status": {str(k): int(v) for k, v in no_school["school_ref_status"].value_counts().items()},
        "status": dict(Counter(r["status"] for r in rows)),
        "family": dict(Counter(r["family"] for r in rows if r["status"] == "eligible")),
        "special": sum(1 for r in rows if r["special"]),
        "la_cy_missing": sorted({r["la"] for r in rows if r["la_cy_missing"]}),
        "la_cy_missing_schools": sum(1 for r in rows if r["la_cy_missing"]),
        "name_from_plasc": sum(1 for r in rows if r["name_source"] == "PLASC 2026"),
        "min_responses": min_responses, "hold_missing_la_cy": hold_missing_la_cy,
        "dataset": str(parquet.name), "plasc": str(plasc.name) if plasc else None,
        "framework": framework.name,
    }
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "register.json").write_text(json.dumps({"summary": summary, "schools": rows}, ensure_ascii=False, indent=1),
                                           encoding="utf-8")
    flat = ["school_id", "slug", "name", "name_source", "la", "la_cy", "region", "partnership", "family", "special",
            "years", "n", "status", "reason"]
    with (out_dir / "register.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(flat)
        for r in rows:
            w.writerow([",".join(map(str, r[k])) if isinstance(r[k], list) else r[k] for k in flat])
    (out_dir / "register_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    return summary


def profile_for(row: dict, release: str, dataset_sha256: str, dataset_file: str) -> dict:
    """The config/schools/<slug>.json the pipeline reads, from a register row —
    the same fields, in the same order, as the V5.1 profiles."""
    fam = row["family"]
    return {
        "profileId": f"standard-{fam}-school",
        "slug": row["slug"],
        "reportVersion": f"2026-school-{release}-bilingual",
        "schoolType": fam,
        "schoolId": row["school_id"],
        "schoolName": row["name"],
        "localAuthority": row["la"],
        "regionalSportPartnership": row["partnership"] or "To be confirmed",
        "schoolStages": row["stages"],
        "teacherSurveyCompleted": "To be confirmed",
        "fieldworkDates": FIELDWORK,
        "fsmBand": "To be confirmed by Sport Wales",
        "surveyYear": 2026,
        "availableYears": row["years"],
        "scopeGroups": row["scope_groups"],
        "genderOptions": [{"key": "all", "label": "All respondents"}, {"key": "boy", "label": "Boys"}, {"key": "girl", "label": "Girls"}],
        "bandedSchool": False,
        "suppressionThreshold": 5,
        "maxAnswerDerivedCohorts": 1,
        "weighting": "none",
        "language": "en",
        "translationReady": True,
        "acceptedStatuses": ["Complete", "Partial"],
        "dataSource": {
            "file": dataset_file, "sha256": dataset_sha256,
            "handover": "SSS2026_pupil_cleansing_handover_v2.3_2026-09-09 (Industryline Research, 9 Sep 2026)",
            "inclusionRule": "completed OR partial with current_page_position >= 48 (analytical_inclusion_flag)",
            "export": "pipeline/cleansed_to_export.py — one school's rows in the prototype's SmartSurvey layout; free text withheld",
        },
    }


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build"); b.add_argument("parquet"); b.add_argument("out_dir")
    b.add_argument("--plasc"); b.add_argument("--min-responses", type=int, default=1,
                   help="owner instruction 22 Sep 2026: every school with an accepted response gets a report (default 1); the view-level rule of five is applied inside the report, not here")
    b.add_argument("--framework", default="config/01_Framework_v2.19.xlsx")
    b.add_argument("--allow-missing-la-cy", action="store_true",
                   help="build schools whose authority has no Welsh row on sheet 53 (default: HELD — owner decision 22 Sep 2026, option A)")
    p = sub.add_parser("profile"); p.add_argument("register"); p.add_argument("school_id"); p.add_argument("out")
    p.add_argument("--release", required=True); p.add_argument("--dataset-sha256", required=True)
    p.add_argument("--dataset-file", default="SSS2026_pupil_stage2_full_cleaned.parquet")
    a = ap.parse_args(argv)
    if a.cmd == "build":
        s = build(Path(a.parquet), Path(a.out_dir), Path(a.plasc) if a.plasc else None, Path(a.framework),
                  a.min_responses, not a.allow_missing_la_cy)
        print(json.dumps(s, ensure_ascii=False, indent=1))
    else:
        reg = json.loads(Path(a.register).read_text(encoding="utf-8"))
        row = next(r for r in reg["schools"] if r["school_id"] == a.school_id)
        Path(a.out).write_text(json.dumps(profile_for(row, a.release, a.dataset_sha256, a.dataset_file),
                                          ensure_ascii=False, indent=2), encoding="utf-8")
        print(a.out)


if __name__ == "__main__":
    main()
